"""Files: metadata, downloads, versions, updates, locks, thumbnails, trash."""

from __future__ import annotations

import hashlib
import io
import re
from datetime import timedelta

import pytest

DELIVERABLES = "Client Engagements/Northwind Logistics/FY2025 Review/Deliverables"


async def file_in(box, folder_path: str, name: str, who: str = "p035") -> dict:
    folder = await box.find(folder_path, who)
    entry = next(e for e in await box.items(folder, who) if e["name"] == name)
    return entry


async def test_file_standard_fields(box):
    entry = await file_in(box, DELIVERABLES, "Management Letter.docx")
    r = await box.get(f"/files/{entry['id']}")
    assert r.status_code == 200
    file = r.json()
    assert list(file)[:4] == ["type", "id", "file_version", "sequence_id"]
    assert file["type"] == "file" and file["name"] == "Management Letter.docx"
    assert file["file_version"]["type"] == "file_version"
    assert file["file_version"]["sha1"] == file["sha1"]
    assert re.fullmatch(r"[0-9a-f]{40}", file["sha1"])
    assert file["size"] > 500
    for key in (
        "description", "path_collection", "created_at", "modified_at", "trashed_at", "purged_at",
        "content_created_at", "content_modified_at", "created_by", "modified_by", "owned_by",
        "shared_link", "parent", "item_status",
    ):  # fmt: skip
        assert key in file, key
    assert file["owned_by"]["login"] == "arjun.mehta@halcyonadvisory.example"
    assert file["created_by"]["login"].endswith("@halcyonadvisory.example")
    assert "extension" not in file and "permissions" not in file
    assert file["parent"]["name"] == "Deliverables"
    assert [e["name"] for e in file["path_collection"]["entries"]][:3] == [
        "All Files", "Client Engagements", "Northwind Logistics",
    ]  # fmt: skip


async def test_file_extended_fields_on_request(box):
    entry = await file_in(box, DELIVERABLES, "Management Letter.docx")
    fields = "extension,version_number,comment_count,permissions,tags,lock,has_collaborations,is_externally_owned,expires_at,collections,metadata.enterprise.clientInfo"
    file = (await box.get(f"/files/{entry['id']}", params={"fields": fields})).json()
    assert file["extension"] == "docx"
    assert int(file["version_number"]) >= 1
    assert file["comment_count"] >= 0 and file["tags"] == [] and file["lock"] is None
    assert file["is_externally_owned"] is False and file["collections"] == []
    assert file["permissions"]["can_download"] and file["permissions"]["can_preview"]
    info = file["metadata"]["enterprise_336533359"]["clientInfo"]
    assert info["clientName"] == "Northwind Logistics" and info["$template"] == "clientInfo"
    assert set(file) >= {"type", "id", "etag", "sequence_id", "name", "sha1", "file_version"}
    assert "path_collection" not in file


async def test_download_redirects_to_boxcloud_and_serves_valid_documents(box):
    entry = await file_in(box, DELIVERABLES, "Management Letter.docx")
    r = await box.get(f"/files/{entry['id']}/content")
    assert r.status_code == 302
    location = r.headers["location"]
    assert re.fullmatch(r"https://dl\.boxcloud\.com/d/1/b1![\w-]+/download", location)
    # The boxcloud URL needs no Authorization header.
    data = await box.client.get(location)
    assert data.status_code == 200
    assert data.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml"
    )
    assert "attachment" in data.headers["content-disposition"]
    assert hashlib.sha1(data.content).hexdigest() == entry["sha1"]
    from docx import Document

    doc = Document(io.BytesIO(data.content))
    assert len(doc.paragraphs) > 1


async def test_download_range_and_head(box):
    entry = await file_in(box, DELIVERABLES, "Draft Financial Statements.pdf")
    location = (await box.get(f"/files/{entry['id']}/content")).headers["location"]
    full = (await box.client.get(location)).content
    part = await box.client.get(location, headers={"Range": "bytes=2-9"})
    assert part.status_code == 206
    assert part.content == full[2:10]
    assert part.headers["content-range"] == f"bytes 2-9/{len(full)}"
    tail = await box.client.get(location, headers={"Range": "bytes=-5"})
    assert tail.content == full[-5:]
    open_ended = await box.client.get(location, headers={"Range": f"bytes={len(full) - 3}-"})
    assert open_ended.content == full[-3:]
    unsatisfiable = await box.client.get(location, headers={"Range": f"bytes={len(full) + 10}-"})
    assert unsatisfiable.status_code == 416
    head = await box.client.head(location)
    assert (
        head.status_code == 200 and int(head.headers["content-length"]) == len(full) and head.content == b""
    )
    assert head.headers["accept-ranges"] == "bytes"


async def test_download_links_expire(box):
    entry = await file_in(box, DELIVERABLES, "Management Letter.docx")
    location = (await box.get(f"/files/{entry['id']}/content")).headers["location"]
    assert (await box.client.get(location)).status_code == 200
    box.hub.clock.advance(timedelta(minutes=20))
    assert (await box.client.get(location)).status_code == 404


async def test_download_access_control(box):
    # Finance staff only have viewer access to the engagement files: they can download but not change.
    entry = await file_in(box, DELIVERABLES, "Management Letter.docx", "p030")
    assert (await box.get(f"/files/{entry['id']}/content", "p030")).status_code == 302
    assert (await box.get(f"/files/{entry['id']}/content", "p044")).status_code == 404
    assert (await box.get(f"/files/{entry['id']}/content", "box-no-such-token")).status_code == 401


async def test_versions_listing_and_download_of_old_versions(box):
    entry = await file_in(box, DELIVERABLES, "Draft Financial Statements.pdf")
    file = (await box.get(f"/files/{entry['id']}", params={"fields": "version_number"})).json()
    current = int(file["version_number"])
    assert current == 4
    versions = (await box.get(f"/files/{entry['id']}/versions")).json()
    assert versions["total_count"] == current - 1
    assert [v["version_number"] for v in versions["entries"]] == ["3", "2", "1"]
    assert file["file_version"]["id"] not in {v["id"] for v in versions["entries"]}
    old = versions["entries"][-1]
    assert old["type"] == "file_version" and re.fullmatch(r"[0-9a-f]{40}", old["sha1"])
    assert old["modified_by"]["type"] == "user" and old["size"] > 0
    location = (await box.get(f"/files/{entry['id']}/content", params={"version": old["id"]})).headers[
        "location"
    ]
    data = (await box.client.get(location)).content
    assert hashlib.sha1(data).hexdigest() == old["sha1"] and len(data) == old["size"]
    assert (
        data
        != (
            await box.client.get((await box.get(f"/files/{entry['id']}/content")).headers["location"])
        ).content
    )
    one = await box.get(f"/files/{entry['id']}/versions/{old['id']}")
    assert one.json()["id"] == old["id"]
    assert (await box.get(f"/files/{entry['id']}/versions/12345")).status_code == 404
    assert (await box.get(f"/files/{entry['id']}/content", params={"version": "12345"})).status_code == 404


async def test_new_versions_promote_delete_and_restore(box):
    parent = await box.find("Technology")
    created = (await box.upload(parent, "plan.txt", b"v1 content")).json()["entries"][0]
    fid = created["id"]
    first_version = created["file_version"]["id"]
    for text in (b"v2 content", b"v3 content"):
        r = await box.client.post(
            f"https://upload.box.com/api/2.0/files/{fid}/content",
            headers=box.auth(),
            files={"file": ("plan.txt", text)},
        )
        assert r.status_code == 201, r.text
    file = (await box.get(f"/files/{fid}", params={"fields": "version_number,size,sha1"})).json()
    assert file["version_number"] == "3" and file["sha1"] == hashlib.sha1(b"v3 content").hexdigest()
    versions = (await box.get(f"/files/{fid}/versions")).json()["entries"]
    assert [v["version_number"] for v in versions] == ["2", "1"]
    promoted = await box.post(
        f"/files/{fid}/versions/current", json={"id": first_version, "type": "file_version"}
    )
    assert promoted.status_code == 201
    assert promoted.json()["version_number"] == "4"
    location = (await box.get(f"/files/{fid}/content")).headers["location"]
    assert (await box.client.get(location)).content == b"v1 content"
    # Deleting a version trashes it; it can be restored.
    victim = versions[0]["id"]
    assert (await box.delete(f"/files/{fid}/versions/{victim}")).status_code == 204
    trashed = (await box.get(f"/files/{fid}/versions/{victim}")).json()
    assert trashed["trashed_at"] and trashed["trashed_by"]["name"] == "Arjun Mehta"
    assert (await box.get(f"/files/{fid}/content", params={"version": victim})).status_code == 404
    restored = await box.put(f"/files/{fid}/versions/{victim}", json={"trashed_at": None})
    assert restored.status_code == 200 and restored.json()["trashed_at"] is None
    assert (await box.get(f"/files/{fid}/content", params={"version": victim})).status_code == 302
    current = (await box.get(f"/files/{fid}")).json()["file_version"]["id"]
    assert (await box.delete(f"/files/{fid}/versions/{current}")).status_code == 403


async def test_update_rename_description_move_and_conflicts(box):
    parent = await box.find("Technology")
    other = await box.mkdir(parent, "Target")
    first = (await box.upload(parent, "a.txt", b"a")).json()["entries"][0]["id"]
    await box.upload(parent, "b.txt", b"b")
    r = await box.put(f"/files/{first}", json={"name": "a-renamed.txt", "description": "Hello"})
    assert r.status_code == 200
    assert r.json()["name"] == "a-renamed.txt" and r.json()["description"] == "Hello"
    assert r.json()["etag"] == "1"
    assert (await box.put(f"/files/{first}", json={"name": "b.txt"})).status_code == 409
    assert (await box.put(f"/files/{first}", json={"description": "x" * 257})).status_code == 400
    assert (await box.put(f"/files/{first}", json={"name": "bad/name"})).status_code == 400
    moved = await box.put(f"/files/{first}", json={"parent": {"id": other}})
    assert moved.json()["parent"]["id"] == other
    assert first in [e["id"] for e in await box.items(other)]
    assert first not in [e["id"] for e in await box.items(parent)]
    assert (
        await box.put(f"/files/{first}", headers={"If-Match": "99"}, json={"name": "z.txt"})
    ).status_code == 412
    tags = await box.put(f"/files/{first}", json={"tags": ["audit", "2026"]})
    assert (await box.get(f"/files/{first}", params={"fields": "tags"})).json()["tags"] == ["audit", "2026"]
    assert tags.status_code == 200


async def test_copy_file(box):
    parent = await box.find("Technology")
    src = (await box.upload(parent, "orig.txt", b"copy me")).json()["entries"][0]
    clash = await box.post(f"/files/{src['id']}/copy", json={"parent": {"id": parent}})
    assert clash.status_code == 409
    assert clash.json()["context_info"]["conflicts"]["id"] == src["id"]
    copy = await box.post(
        f"/files/{src['id']}/copy", json={"parent": {"id": parent}, "name": "orig copy.txt"}
    )
    assert copy.status_code == 201
    body = copy.json()
    assert body["id"] != src["id"] and body["sha1"] == src["sha1"] and body["etag"] == "0"
    location = (await box.get(f"/files/{body['id']}/content")).headers["location"]
    assert (await box.client.get(location)).content == b"copy me"


async def test_trash_restore_and_purge_a_file(box):
    parent = await box.find("Technology")
    fid = (await box.upload(parent, "temp.txt", b"temp")).json()["entries"][0]["id"]
    assert (await box.delete(f"/files/{fid}")).status_code == 204
    gone = await box.get(f"/files/{fid}")
    assert gone.status_code == 404 and gone.json()["code"] == "trashed"
    trashed = (await box.get(f"/files/{fid}/trash")).json()
    assert trashed["item_status"] == "trashed"
    assert fid in [e["id"] for e in (await box.get("/folders/trash/items")).json()["entries"]]
    restored = await box.post(f"/files/{fid}", json={})
    assert restored.status_code == 201 and restored.json()["item_status"] == "active"
    assert (await box.get(f"/files/{fid}")).status_code == 200
    await box.delete(f"/files/{fid}")
    assert (await box.delete(f"/files/{fid}/trash")).status_code == 204
    assert (await box.get(f"/files/{fid}/trash")).status_code == 404
    assert (await box.post(f"/files/{fid}", json={})).status_code == 404


async def test_file_locks(box):
    parent = await box.find("Client Engagements/Northwind Logistics/Correspondence")
    fid = (await box.upload(parent, "locked.txt", b"v1")).json()["entries"][0]["id"]
    expires = (box.hub.clock.now() + timedelta(hours=1)).isoformat()
    # p004 (editor via Client Engagements) can lock; the lock is visible with ?fields=lock.
    locked = await box.put(
        f"/files/{fid}",
        "p004",
        json={"lock": {"access": "lock", "expires_at": expires, "is_download_prevented": True}},
    )
    assert locked.status_code == 200
    lock = (await box.get(f"/files/{fid}", params={"fields": "lock"})).json()["lock"]
    assert lock["type"] == "lock" and lock["created_by"]["name"] == "Daniel Kowalski"
    assert lock["is_download_prevented"] is True and lock["expired_at"]
    # Other users cannot modify, replace or download while it is held.
    blocked = await box.put(f"/files/{fid}", "p019", json={"name": "other.txt"})
    assert blocked.status_code == 403 and blocked.json()["code"] == "access_denied_item_locked"
    version = await box.client.post(
        f"https://upload.box.com/api/2.0/files/{fid}/content",
        headers=box.auth("p019"),
        files={"file": ("x", b"y")},
    )
    assert version.status_code == 403
    assert (await box.delete(f"/files/{fid}", "p019")).status_code == 403
    assert (await box.get(f"/files/{fid}/content", "p019")).status_code == 403
    assert (await box.put(f"/files/{fid}", "p019", json={"lock": None})).status_code == 403
    assert (await box.get(f"/files/{fid}/content", "p004")).status_code == 302
    assert (await box.put(f"/files/{fid}", "p004", json={"description": "mine"})).status_code == 200
    # Locks expire by themselves.
    box.hub.clock.advance(timedelta(hours=2))
    assert (await box.put(f"/files/{fid}", "p019", json={"description": "now ok"})).status_code == 200
    # Explicit unlock.
    await box.put(f"/files/{fid}", "p004", json={"lock": {"access": "lock"}})
    assert (await box.put(f"/files/{fid}", "p004", json={"lock": None})).status_code == 200
    assert (await box.get(f"/files/{fid}", params={"fields": "lock"})).json()["lock"] is None


async def test_thumbnails(box):
    entry = await file_in(box, DELIVERABLES, "Management Letter.docx")
    r = await box.get(f"/files/{entry['id']}/thumbnail.png", params={"min_height": 128, "min_width": 128})
    assert r.status_code == 200 and r.headers["content-type"] == "image/png"
    assert r.content.startswith(b"\x89PNG")
    default = await box.get(f"/files/{entry['id']}/thumbnail.png")
    assert default.status_code == 200
    assert (
        await box.get(f"/files/{entry['id']}/thumbnail.png", params={"min_height": 100})
    ).status_code == 400
    assert (await box.get(f"/files/{entry['id']}/thumbnail.gif")).status_code == 404
    assert (await box.get(f"/files/{entry['id']}/thumbnail.png", "p044")).status_code == 404


async def test_missing_and_malformed_file_ids(box):
    for path in ("/files/1", "/files/not-a-number", "/files/0"):
        r = await box.get(path)
        assert r.status_code == 404 and r.json()["type"] == "error", path
    assert (await box.get("/files/1/content")).status_code == 404
    assert (await box.put("/files/1", json={})).status_code == 404
    assert (await box.delete("/files/1")).status_code == 404


@pytest.mark.parametrize("kind", ["xlsx", "pdf", "pptx"])
async def test_seeded_binary_documents_open_in_real_parsers(box, kind):
    # Walk the whole tree once and open the first seeded document of each kind.
    stack = ["0"]
    target = None
    while stack and target is None:
        for entry in await box.items(stack.pop(), fields="name"):
            if entry["type"] == "folder":
                stack.append(entry["id"])
            elif entry["name"].endswith(f".{kind}"):
                target = entry
                break
    assert target is not None
    location = (await box.get(f"/files/{target['id']}/content")).headers["location"]
    data = (await box.client.get(location)).content
    if kind == "xlsx":
        import openpyxl

        assert openpyxl.load_workbook(io.BytesIO(data)).sheetnames
    elif kind == "pdf":
        import pypdf

        assert len(pypdf.PdfReader(io.BytesIO(data)).pages) >= 1
    else:
        import pptx

        assert len(pptx.Presentation(io.BytesIO(data)).slides) >= 1


async def test_prefix_routing_generates_urls_on_the_same_base(make_hub):
    """Behind ``http://localhost:44080/box`` every absolute URL stays on that base."""
    import httpx

    hub = make_hub("box")
    transport = httpx.ASGITransport(app=hub.gateway)
    auth = {"Authorization": "Bearer box-dev-p035"}
    async with httpx.AsyncClient(
        transport=transport, base_url="http://localhost:44080/box", follow_redirects=False
    ) as c:
        root = (await c.get("/2.0/folders/0/items", headers=auth)).json()["entries"]
        tech = next(e["id"] for e in root if e["name"] == "Technology")
        up = await c.post(
            "/api/2.0/files/content",
            headers=auth,
            files={
                "attributes": (None, f'{{"name": "p.txt", "parent": {{"id": "{tech}"}}}}'),
                "file": ("p.txt", b"prefix"),
            },
        )
        assert up.status_code == 201
        fid = up.json()["entries"][0]["id"]
        redirect = await c.get(f"/2.0/files/{fid}/content", headers=auth)
        assert redirect.headers["location"].startswith("http://localhost:44080/box/d/1/")
        assert (await c.get(redirect.headers["location"])).content == b"prefix"
        session = await c.post(
            "/api/2.0/files/upload_sessions",
            headers=auth,
            json={"folder_id": tech, "file_size": 5, "file_name": "s.bin"},
        )
        assert session.json()["session_endpoints"]["commit"].startswith(
            "http://localhost:44080/box/api/2.0/files/upload_sessions/"
        )
        link = await c.put(
            f"/2.0/files/{fid}",
            headers=auth,
            params={"fields": "shared_link"},
            json={"shared_link": {"access": "open"}},
        )
        assert link.json()["shared_link"]["url"].startswith("http://localhost:44080/box/s/")
        longpoll = await c.request("OPTIONS", "/2.0/events", headers=auth)
        assert longpoll.json()["entries"][0]["url"].startswith(
            "http://localhost:44080/box/subscribe?channel="
        )
