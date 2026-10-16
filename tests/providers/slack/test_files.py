"""Files: the external upload flow, legacy files.upload, listing, downloads and permissions."""

from __future__ import annotations

import io
from typing import TYPE_CHECKING, Any

import httpx

if TYPE_CHECKING:
    from .conftest import Slack

SLACK = "https://slack.com"


async def upload_v2(
    slack: Slack, token: str | None, name: str, data: bytes, *, channel: str | None = None, **extra: Any
) -> dict[str, Any]:
    start = await slack.ok("files.getUploadURLExternal", token, filename=name, length=len(data))
    sent = await slack.client.post(start["upload_url"], content=data)
    assert sent.status_code == 200 and sent.text == f"OK - {len(data)}"
    return await slack.ok(
        "files.completeUploadExternal",
        token,
        files=[{"id": start["file_id"], "title": extra.pop("title", name)}],
        channel_id=channel,
        **extra,
    )


async def test_external_upload_flow(slack: Slack) -> None:
    channel = slack.channel_id("general")
    payload = b"quarterly numbers\n1,2,3\n"
    start = await slack.ok("files.getUploadURLExternal", filename="numbers.csv", length=len(payload))
    assert start["upload_url"].startswith("https://files.slack.com/upload/v1/")
    assert start["file_id"].startswith("F")
    sent = await slack.client.post(start["upload_url"], content=payload)
    assert (sent.status_code, sent.text) == (200, f"OK - {len(payload)}")
    done = await slack.ok(
        "files.completeUploadExternal",
        files=[{"id": start["file_id"], "title": "Numbers"}],
        channel_id=channel,
        initial_comment="Here you go",
    )
    (record,) = done["files"]
    assert (
        record["id"] == start["file_id"] and record["title"] == "Numbers" and record["name"] == "numbers.csv"
    )
    assert record["size"] == len(payload) and record["mimetype"] == "text/csv" and record["filetype"] == "csv"
    assert record["url_private"].startswith(
        f"https://files.slack.com/files-pri/{slack.ctx.world.org.slack_team_id}-{record['id']}/"
    )
    assert record["url_private_download"].endswith("/download/numbers.csv")
    assert record["permalink"].startswith("https://halcyonadvisory.slack.com/files/")
    assert (
        record["channels"] == [channel]
        and record["user"] == slack.ctx.store.collection("app_config").get("app")["bot_user_id"]
    )
    shown = (await slack.history(channel, limit=1))[0]
    assert shown["subtype"] == "file_share" and shown["text"] == "Here you go"
    assert shown["upload"] is True and shown["files"][0]["id"] == record["id"]
    info = (await slack.ok("files.info", file=record["id"]))["file"]
    assert info["shares"]["public"][channel][0]["ts"] == shown["ts"]
    token = await slack.bot()
    download = await slack.client.get(
        record["url_private_download"], headers={"Authorization": f"Bearer {token}"}
    )
    assert download.content == payload and download.headers["content-type"].startswith("text/csv")
    assert "attachment" in download.headers["content-disposition"]
    view = await slack.client.get(record["url_private"], headers={"Authorization": f"Bearer {token}"})
    assert view.content == payload and "content-disposition" not in view.headers


async def test_upload_in_a_thread_and_multipart_upload_url(slack: Slack) -> None:
    channel = slack.channel_id("general")
    parent = (await slack.ok("chat.postMessage", channel=channel, text="thread"))["ts"]
    start = await slack.ok("files.getUploadURLExternal", filename="notes.txt", length=5)
    sent = await slack.client.post(start["upload_url"], files={"file": ("notes.txt", b"hello", "text/plain")})
    assert sent.text == "OK - 5"
    done = await slack.ok(
        "files.completeUploadExternal",
        files=[{"id": start["file_id"]}],
        channel_id=channel,
        thread_ts=parent,
        initial_comment="in thread",
    )
    assert done["files"][0]["title"] == "notes.txt"
    replies = (await slack.ok("conversations.replies", channel=channel, ts=parent))["messages"]
    assert replies[1]["files"][0]["id"] == start["file_id"] and replies[1]["thread_ts"] == parent


async def test_complete_upload_errors(slack: Slack) -> None:
    channel = slack.channel_id("general")
    start = await slack.ok("files.getUploadURLExternal", filename="x.txt", length=3)
    assert (
        await slack.err("files.completeUploadExternal", files=[{"id": start["file_id"]}], channel_id=channel)
        == "file_not_found"
    )
    assert await slack.err("files.completeUploadExternal", files="nope") == "invalid_arguments"
    assert await slack.err("files.completeUploadExternal", files=[]) == "invalid_arguments"
    assert await slack.err("files.getUploadURLExternal", filename="x.txt") == "invalid_arguments"
    assert await slack.err("files.getUploadURLExternal", length=3) == "invalid_arguments"
    assert await slack.client.post("https://files.slack.com/upload/v1/unknown", content=b"x") is not None
    missing = await slack.client.post("https://files.slack.com/upload/v1/unknown", content=b"x")
    assert missing.status_code == 404
    await slack.client.post(start["upload_url"], content=b"abc")
    assert (
        await slack.err(
            "files.completeUploadExternal",
            files=[{"id": start["file_id"]}],
            channel_id=slack.channel_id("social"),
        )
        == "not_in_channel"
    )
    other = await slack.user_token("p018")
    assert (
        await slack.err(
            "files.completeUploadExternal", other, files=[{"id": start["file_id"]}], channel_id=channel
        )
        == "file_not_found"
    )
    private = (await slack.ok("files.completeUploadExternal", files=[{"id": start["file_id"]}]))["files"][0]
    assert private["channels"] == [] and private["shares"] == {}
    assert (await slack.ok("files.info", file=private["id"]))["file"]["id"] == private["id"]
    assert await slack.err("files.info", other, file=private["id"]) == "file_not_found"


async def test_legacy_files_upload(slack: Slack) -> None:
    channel = slack.channel_id("general")
    token = await slack.bot()
    r = await slack.client.post(
        f"{SLACK}/api/files.upload",
        headers={"Authorization": f"Bearer {token}"},
        data={"channels": channel, "title": "Report", "initial_comment": "see attached"},
        files={"file": ("report.txt", b"legacy upload", "text/plain")},
    )
    body = r.json()
    assert body["ok"] is True and body["file"]["name"] == "report.txt" and body["file"]["title"] == "Report"
    assert body["file"]["channels"] == [channel] and body["file"]["mimetype"] == "text/plain"
    content = await slack.ok(
        "files.upload", None, content="inline body", filename="inline.md", channels=channel
    )
    assert content["file"]["filetype"] == "markdown" and content["file"]["size"] == len("inline body")
    assert await slack.err("files.upload", None, channels=channel) == "no_file_data"
    assert await slack.err("files.upload", None, content="x", channels="C0NOTREALXX") == "channel_not_found"
    assert (
        await slack.err("files.upload", None, content="x", channels=slack.channel_id("social"))
        == "not_in_channel"
    )
    private_only = await slack.ok("files.upload", None, content="just me", filename="me.txt")
    assert private_only["file"]["channels"] == []


async def test_download_requires_a_token_with_files_read(slack: Slack) -> None:
    channel = slack.channel_id("general")
    done = await upload_v2(slack, None, "secret.txt", b"classified", channel=channel)
    url = done["files"][0]["url_private"]
    anonymous = await slack.client.get(url)
    assert anonymous.status_code == 302 and "halcyonadvisory.slack.com" in anonymous.headers["location"]
    garbage = await slack.client.get(url, headers={"Authorization": "Bearer xoxb-garbage"})
    assert garbage.status_code == 302
    narrow = await slack.mint(scopes=["chat:write"])
    forbidden = await slack.client.get(url, headers={"Authorization": f"Bearer {narrow}"})
    assert forbidden.status_code == 403
    wrong = await slack.client.get(
        "https://files.slack.com/files-pri/TOTHERTEAM-F123/x.txt",
        headers={"Authorization": f"Bearer {await slack.bot()}"},
    )
    assert wrong.status_code == 404
    # A user outside a private conversation cannot fetch its files.
    private = slack.channel_id("partners")
    eleanor = await slack.user_token("p001")
    shared = await upload_v2(slack, eleanor, "partners.txt", b"p", channel=private)
    maya = await slack.mint(user="p020")
    blocked = await slack.client.get(
        shared["files"][0]["url_private"], headers={"Authorization": f"Bearer {maya}"}
    )
    assert blocked.status_code == 404
    allowed = await slack.client.get(
        shared["files"][0]["url_private"], headers={"Authorization": f"Bearer {eleanor}"}
    )
    assert allowed.content == b"p"


async def test_seeded_documents_download_as_real_files(slack: Slack) -> None:
    token = await slack.user_token("p001")
    files = (await slack.ok("files.list", token, count=100))["files"]
    by_type = {}
    for record in files:
        by_type.setdefault(record["filetype"], record)
    assert by_type
    parsed = 0
    for kind, record in by_type.items():
        data = (
            await slack.client.get(
                record["url_private_download"], headers={"Authorization": f"Bearer {token}"}
            )
        ).content
        assert len(data) == record["size"] > 100
        if kind == "docx":
            import docx

            assert docx.Document(io.BytesIO(data)).paragraphs
            parsed += 1
        elif kind == "xlsx":
            import openpyxl

            assert openpyxl.load_workbook(io.BytesIO(data)).sheetnames
            parsed += 1
        elif kind == "pdf":
            import pypdf

            assert len(pypdf.PdfReader(io.BytesIO(data)).pages) >= 1
            parsed += 1
    assert parsed >= 1
    assert all(r["mimetype"] and r["pretty_type"] for r in files)


async def test_files_list_filters_and_paging(slack: Slack) -> None:
    eleanor = await slack.user_token("p001")
    general, social = slack.channel_id("general"), slack.channel_id("it-help")
    await upload_v2(slack, eleanor, "a.png", b"\x89PNG", channel=general)
    await upload_v2(slack, eleanor, "b.pdf", b"%PDF-1.4", channel=social)
    await upload_v2(slack, eleanor, "c.txt", b"text", channel=general)
    mine = await slack.ok("files.list", eleanor, user=slack.uid("p001"))
    assert [f["name"] for f in mine["files"]] == ["c.txt", "b.pdf", "a.png"]
    assert mine["paging"] == {"count": 100, "total": 3, "page": 1, "pages": 1}
    in_general = await slack.ok("files.list", eleanor, user=slack.uid("p001"), channel=general)
    assert {f["name"] for f in in_general["files"]} == {"a.png", "c.txt"}
    images = await slack.ok("files.list", eleanor, user=slack.uid("p001"), types="images")
    assert [f["name"] for f in images["files"]] == ["a.png"]
    assert images["files"][0]["thumb_64"].startswith("https://files.slack.com/files-tmb/")
    pdfs = await slack.ok("files.list", eleanor, user=slack.uid("p001"), types="pdfs,images")
    assert {f["name"] for f in pdfs["files"]} == {"a.png", "b.pdf"}
    page = await slack.ok("files.list", eleanor, user=slack.uid("p001"), count=2, page=2)
    assert [f["name"] for f in page["files"]] == ["a.png"] and page["paging"]["pages"] == 2
    future = await slack.ok("files.list", eleanor, user=slack.uid("p001"), ts_from=4_000_000_000)
    assert future["files"] == []


async def test_delete_files(slack: Slack) -> None:
    channel = slack.channel_id("general")
    eleanor = await slack.user_token("p001")
    done = await upload_v2(slack, eleanor, "gone.txt", b"bye", channel=channel)
    file_id = done["files"][0]["id"]
    maya = await slack.user_token("p018")
    assert await slack.err("files.delete", maya, file=file_id) == "cant_delete_file"
    admin = await slack.user_token("p035")
    assert await slack.ok("files.delete", admin, file=file_id) == {"ok": True}
    assert await slack.err("files.info", None, file=file_id) == "file_deleted"
    assert await slack.err("files.delete", admin, file=file_id) == "file_deleted"
    assert await slack.err("files.delete", admin, file="FNOSUCHFILE") == "file_not_found"
    shown = (await slack.history(channel, limit=1))[0]
    assert shown["files"] == []
    token = await slack.bot()
    assert (
        await slack.client.get(done["files"][0]["url_private"], headers={"Authorization": f"Bearer {token}"})
    ).status_code == 404


async def test_public_urls_and_file_reactions(slack: Slack) -> None:
    eleanor = await slack.user_token("p001")
    done = await upload_v2(slack, eleanor, "share.txt", b"s", channel=slack.channel_id("general"))
    file_id = done["files"][0]["id"]
    shared = (await slack.ok("files.sharedPublicURL", eleanor, file=file_id))["file"]
    assert shared["public_url_shared"] is True and shared["permalink_public"].startswith(
        "https://slack-files.com/T"
    )
    assert (await slack.ok("files.revokePublicURL", eleanor, file=file_id))["file"][
        "public_url_shared"
    ] is False
    assert (await slack.ok("files.sharedPublicURL", await slack.user_token("p018"), file=file_id))["file"][
        "public_url_shared"
    ] is True
    await slack.ok("reactions.add", eleanor, file=file_id, name="tada")
    got = (await slack.ok("reactions.get", eleanor, file=file_id))["file"]
    assert got["reactions"][0]["name"] == "tada"
    assert await slack.err("reactions.add", eleanor, file="FNOSUCHFILE", name="tada") == "file_not_found"


async def test_unauthenticated_file_requests_do_not_leak_on_other_hosts(slack: Slack) -> None:
    r = await slack.client.get(f"{SLACK}/files-pri/T0-F0/x.txt")
    assert r.status_code == 302
    assert isinstance(r, httpx.Response)


async def test_legacy_files_upload_can_be_switched_off() -> None:
    from emulators.core.hub import EmulatorHub

    hub = EmulatorHub(
        providers=["slack"], anchor="2026-10-16", settings={"slack": {"legacy_files_upload": "false"}}
    )
    token = hub.context("slack").store.collection("_grants")
    bot = next(r["token"] for r in token if r["token"].startswith("xoxb-"))
    async with httpx.AsyncClient(transport=hub.async_transport(), base_url=SLACK) as client:
        r = await client.post(
            "/api/files.upload", headers={"Authorization": f"Bearer {bot}"}, data={"content": "x"}
        )
        assert r.json() == {"ok": False, "error": "method_deprecated"}
