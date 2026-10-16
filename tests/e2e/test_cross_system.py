"""Reconcile the shared world through public vendor APIs, never provider stores.

Each test uses the same cached, resettable sixteen-provider hub from make_client.
The immutable world is an independent oracle; matching two incomplete projections
alone would miss dropped people, invoices or documents.
"""

from __future__ import annotations

import hashlib
import hmac
import io
import json
from collections.abc import AsyncIterator, Callable
from dataclasses import dataclass
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from html.parser import HTMLParser
from itertools import groupby
from typing import Any
from urllib.parse import parse_qs, urlsplit

import httpx
import pytest
from docx import Document as DocxDocument
from fastapi import FastAPI, Request
from fastapi.responses import PlainTextResponse
from openpyxl import load_workbook
from pypdf import PdfReader

from emulators.testing.helpers import bearer, mint
from emulators.world import media
from emulators.world.model import Document, Meeting, World

pytestmark = pytest.mark.e2e
PROVIDERS = (
    "microsoft",
    "google",
    "box",
    "dropbox",
    "salesforce",
    "zoom",
    "slack",
    "twilio",
    "meta",
    "stripe",
    "quickbooks",
    "xero",
    "hubspot",
    "webex",
    "notion",
    "atlassian",
)
GRAPH = "https://graph.microsoft.com/v1.0"
GOOGLE = "https://www.googleapis.com"
SF = "https://halcyonadvisory.my.salesforce.com"
HUBSPOT = "https://api.hubapi.com"
BOX = "https://api.box.com"
DROPBOX = "https://api.dropboxapi.com"
ZOOM = "https://api.zoom.us"
WEBEX = "https://webexapis.com"
STRIPE = "https://api.stripe.com"


@dataclass
class APIs:
    client: httpx.AsyncClient
    world: World

    async def auth(self, host: str, **payload: Any) -> dict[str, str]:
        return bearer(await mint(self.client, host, **payload))

    async def request(self, method: str, url: str, **kwargs: Any) -> dict[str, Any]:
        response = await self.client.request(method, url, **kwargs)
        assert response.is_success, f"{method} {url}: {response.status_code} {response.text[:1000]}"
        result = response.json()
        assert result.get("ok", True), result
        return result

    async def get(self, url: str, **kwargs: Any) -> dict[str, Any]:
        return await self.request("GET", url, **kwargs)

    async def post(self, url: str, **kwargs: Any) -> dict[str, Any]:
        return await self.request("POST", url, **kwargs)

    async def pages(
        self,
        url: str,
        key: str,
        *,
        headers: dict[str, str],
        params: dict[str, Any] | None = None,
        next_key: str = "nextPageToken",
        token_param: str = "pageToken",
    ) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = []
        query = dict(params or {})
        seen: set[str] = set()
        while True:
            page = await self.get(url, headers=headers, params=query)
            rows.extend(page.get(key, []))
            token = page.get(next_key)
            if not token:
                return rows
            assert token not in seen, f"Repeated cursor on {url}"
            seen.add(token)
            if next_key == "@odata.nextLink":
                url, query = token, None
            else:
                query[token_param] = token

    async def salesforce(self, query: str) -> list[dict[str, Any]]:
        auth = await self.auth(SF, user="p035")
        page = await self.get(SF + "/services/data/v62.0/query", headers=auth, params={"q": query})
        rows = page["records"]
        while not page["done"]:
            page = await self.get(SF + page["nextRecordsUrl"], headers=auth)
            rows.extend(page["records"])
        return rows

    async def hubspot(self, kind: str, properties: str) -> list[dict[str, Any]]:
        headers = await self.auth(HUBSPOT)
        url = f"{HUBSPOT}/crm/v3/objects/{kind}"
        params: dict[str, Any] = {"limit": 100, "properties": properties}
        rows: list[dict[str, Any]] = []
        while True:
            page = await self.get(url, headers=headers, params=params)
            rows.extend(page["results"])
            after = page.get("paging", {}).get("next", {}).get("after")
            if after is None:
                return rows
            params["after"] = after

    async def qbo(self, entity: str) -> list[dict[str, Any]]:
        host = "https://quickbooks.api.intuit.com"
        page = await self.get(
            f"{host}/v3/company/{self.world.org.quickbooks_realm_id}/query",
            headers={**await self.auth(host), "Accept": "application/json"},
            params={"query": f"SELECT * FROM {entity} MAXRESULTS 1000"},
        )
        return page["QueryResponse"].get(entity, [])

    async def xero(self, entity: str) -> list[dict[str, Any]]:
        host = "https://api.xero.com"
        headers = await self.auth(host, user="p035")
        connections = await self.get_list(host + "/connections", headers=headers)
        tenant = next(c for c in connections if c["tenantName"] == self.world.org.name)
        headers.update({"xero-tenant-id": tenant["tenantId"], "Accept": "application/json"})
        return (await self.get(host + "/api.xro/2.0/" + entity, headers=headers))[entity]

    async def get_list(self, url: str, **kwargs: Any) -> list[dict[str, Any]]:
        response = await self.client.get(url, **kwargs)
        assert response.is_success, response.text
        return response.json()


@pytest.fixture
async def apis(make_client: Callable[..., Any]) -> AsyncIterator[APIs]:
    async with make_client(*PROVIDERS) as client:
        world = client.hub.world
        assert world.anchor.isoformat() == "2026-10-16T00:00:00+00:00"
        yield APIs(client, world)


async def directories(apis: APIs) -> dict[str, dict[str, tuple[str, bool, str]]]:
    """Normalize only identity/status fields, preserving vendor IDs for relationships."""
    result: dict[str, dict[str, tuple[str, bool, str]]] = {}
    users = await apis.pages(
        GRAPH + "/users",
        "value",
        headers=await apis.auth("https://graph.microsoft.com"),
        params={"$top": 19, "$select": "id,mail,displayName,accountEnabled,userType"},
        next_key="@odata.nextLink",
    )
    result["microsoft"] = {u["mail"]: (u["displayName"], u["accountEnabled"], u["id"]) for u in users}
    users = await apis.pages(
        "https://admin.googleapis.com/admin/directory/v1/users",
        "users",
        headers=await apis.auth(GOOGLE, user="p035"),
        params={"customer": "my_customer", "maxResults": 19},
    )
    result["google"] = {
        u["primaryEmail"]: (u["name"]["fullName"], not u["suspended"], u["id"]) for u in users
    }
    users = await apis.pages(
        "https://slack.com/api/users.list",
        "members",
        headers=await apis.auth("https://slack.com", user="p035", scopes=["users:read", "users:read.email"]),
        params={"limit": 200},
    )
    result["slack"] = {
        u["profile"]["email"]: (u["real_name"], not u["deleted"], u["id"])
        for u in users
        if u.get("profile", {}).get("email")
    }
    headers = await apis.auth(ZOOM, user="p035")
    users = []
    for status in ("active", "inactive"):
        users.extend(
            await apis.pages(
                ZOOM + "/v2/users",
                "users",
                headers=headers,
                params={"status": status, "page_size": 19},
                next_key="next_page_token",
                token_param="next_page_token",
            )
        )
    result["zoom"] = {
        u["email"]: (
            u.get("display_name") or f"{u['first_name']} {u['last_name']}",
            u["status"] == "active",
            u["id"],
        )
        for u in users
    }
    users = await apis.salesforce("SELECT Id, Name, Email, IsActive, ManagerId FROM User")
    result["salesforce"] = {u["Email"]: (u["Name"], u["IsActive"], u["Id"]) for u in users}
    headers = await apis.auth(BOX, user="p035")
    users = []
    offset = 0
    while True:
        page = await apis.get(BOX + "/2.0/users", headers=headers, params={"limit": 19, "offset": offset})
        users.extend(page["entries"])
        offset += page["limit"]
        if offset >= page["total_count"]:
            break
    result["box"] = {u["login"]: (u["name"], u["status"] == "active", u["id"]) for u in users}
    headers = await apis.auth(DROPBOX, user="p035", team=True)
    page = await apis.post(DROPBOX + "/2/team/members/list", headers=headers, json={"limit": 19})
    users = page["members"]
    while page["has_more"]:
        page = await apis.post(
            DROPBOX + "/2/team/members/list/continue", headers=headers, json={"cursor": page["cursor"]}
        )
        users.extend(page["members"])
    result["dropbox"] = {
        u["profile"]["email"]: (
            u["profile"]["name"]["display_name"],
            u["profile"]["status"][".tag"] == "active",
            u["profile"]["team_member_id"],
        )
        for u in users
    }
    users = await apis.get(WEBEX + "/v1/people", headers=await apis.auth(WEBEX), params={"max": 1000})
    # Webex status is presence, not login state. License removal marks departed staff.
    result["webex"] = {
        u["emails"][0]: (u["displayName"], bool(u["licenses"]), u["id"]) for u in users["items"]
    }
    return result


async def test_every_employee_has_the_same_directory_identity(apis: APIs) -> None:
    systems = await directories(apis)
    for provider, users in systems.items():
        for person in apis.world.employees:
            assert person.email in users, (provider, person.email)
            name, active, _ = users[person.email]
            assert name == person.display_name, (provider, person.email, name)
            assert active == person.active, (provider, person.email, active)
    owners = []
    for archived in ("false", "true"):
        owners.extend(
            await apis.pages(
                HUBSPOT + "/crm/v3/owners",
                "results",
                headers=await apis.auth(HUBSPOT),
                params={"limit": 100, "archived": archived},
            )
        )
    assert owners
    by_email = {o["email"]: o for o in owners}
    required = {
        p.email
        for p in apis.world.employees
        if p.is_admin
        or p.is_partner
        or p.department_id == "growth"
        or "Manager" in p.job_title
        or any(d.owner_id == p.id for d in apis.world.deals)
        or any(c.owner_id == p.id for c in apis.world.cases)
    }
    assert required == by_email.keys()
    for email, owner in by_email.items():
        person = apis.world.person_by_email(email)
        assert person is not None and not person.is_external
        assert f"{owner['firstName']} {owner['lastName']}" == person.display_name
        assert owner["archived"] == (not person.active)
        assert bool(owner["userId"]) == person.active


async def test_departed_staff_cannot_authenticate(apis: APIs) -> None:
    departed = [p for p in apis.world.employees if not p.active]
    assert departed
    for person in departed:
        for host in (ZOOM, WEBEX):
            response = await apis.client.post(host + "/_emulator/tokens", json={"user": person.email})
            if response.is_success:
                paths = {ZOOM: "/v2/users/me", WEBEX: "/v1/people/me"}
                response = await apis.client.get(
                    host + paths[host], headers=bearer(response.json()["access_token"])
                )
            assert response.status_code in (400, 401, 403), (host, person.email, response.text)
        response = await apis.client.get(
            "https://accounts.google.com/o/oauth2/v2/auth",
            params={
                "client_id": "cross-system-client",
                "redirect_uri": "https://consumer.example/callback",
                "response_type": "code",
                "scope": "openid email",
                "login_hint": person.email,
            },
        )
        assert response.status_code == 302
        assert parse_qs(urlsplit(response.headers["location"]).query)["error"] == ["access_denied"]


async def test_manager_chain_agrees_across_directories(apis: APIs) -> None:
    google_users = await apis.pages(
        "https://admin.googleapis.com/admin/directory/v1/users",
        "users",
        headers=await apis.auth(GOOGLE, user="p035"),
        params={"customer": "my_customer", "maxResults": 500},
    )
    google = {u["primaryEmail"]: u for u in google_users}
    sf_users = await apis.salesforce("SELECT Id, Email, ManagerId FROM User")
    sf = {u["Email"]: u for u in sf_users}
    sf_email = {u["Id"]: u["Email"] for u in sf_users}
    auth = await apis.auth("https://graph.microsoft.com")
    for person in apis.world.employees:
        expected = apis.world.person(person.manager_id).email if person.manager_id else None
        response = await apis.client.get(GRAPH + f"/users/{person.email}/manager", headers=auth)
        if expected is None:
            assert response.status_code == 404
        else:
            assert response.status_code == 200, response.text
            assert response.json()["mail"] == expected
        relations = [r["value"] for r in google[person.email].get("relations", []) if r["type"] == "manager"]
        assert relations == ([expected] if expected else [])
        assert sf_email.get(sf[person.email]["ManagerId"]) == expected


async def test_active_clients_exist_in_every_crm_and_ledger(apis: APIs) -> None:
    expected = {c.name for c in apis.world.clients if c.status == "Active"}
    salesforce = {r["Name"] for r in await apis.salesforce("SELECT Id, Name FROM Account")}
    hubspot = {r["properties"]["name"] for r in await apis.hubspot("companies", "name")}
    qbo = {r["DisplayName"] for r in await apis.qbo("Customer") if r["Active"]}
    xero = {r["Name"] for r in await apis.xero("Contacts") if r["ContactStatus"] == "ACTIVE"}
    assert expected
    for provider, names in {
        "salesforce": salesforce,
        "hubspot": hubspot,
        "quickbooks": qbo,
        "xero": xero,
    }.items():
        assert expected <= names, (provider, expected - names)
    customers = await apis.pages(
        STRIPE + "/v1/customers",
        "data",
        headers={"Authorization": "Bearer sk_test_emulator"},
        params={"limit": 100},
    )
    expected_cards = {c.id for c in apis.world.clients if c.pays_by_card}
    actual_cards = {c["metadata"]["client_id"] for c in customers}
    assert actual_cards == expected_cards
    for customer in customers:
        client = apis.world.client(customer["metadata"]["client_id"])
        assert customer["name"] == client.legal_name
        assert customer["email"] == apis.world.person(client.primary_contact_id).email


async def test_invoice_numbers_and_totals_reconcile(apis: APIs) -> None:
    qbo = {r["DocNumber"]: r for r in await apis.qbo("Invoice")}
    xero = {r["InvoiceNumber"]: r for r in await apis.xero("Invoices") if r["Type"] == "ACCREC"}
    assert set(qbo) == set(xero) == {i.number for i in apis.world.invoices}
    for invoice in apis.world.invoices:
        expected = Decimal(str(invoice.total))
        # QBO voids zero the transaction; Xero keeps its historical face value.
        qbo_expected = Decimal(0) if invoice.status == "Void" else expected
        assert Decimal(str(qbo[invoice.number]["TotalAmt"])) == qbo_expected, invoice.number
        assert Decimal(str(xero[invoice.number]["Total"])) == expected, invoice.number
        balance = Decimal(0) if invoice.status == "Void" else Decimal(str(invoice.balance))
        assert Decimal(str(qbo[invoice.number]["Balance"])) == balance
        assert Decimal(str(xero[invoice.number]["AmountDue"])) == balance


async def test_paid_card_invoices_have_matching_stripe_charges(apis: APIs) -> None:
    headers = {"Authorization": "Bearer sk_test_emulator"}
    customers = await apis.pages(STRIPE + "/v1/customers", "data", headers=headers, params={"limit": 100})
    client_ids = {c["id"]: c["metadata"]["client_id"] for c in customers}
    charges: list[dict[str, Any]] = []
    params: dict[str, Any] = {"limit": 100}
    while True:
        page = await apis.get(STRIPE + "/v1/charges", headers=headers, params=params)
        charges.extend(page["data"])
        if not page["has_more"]:
            break
        params["starting_after"] = page["data"][-1]["id"]
    by_payment = {
        c["metadata"]["world_payment_id"]: c for c in charges if "world_payment_id" in c["metadata"]
    }
    payments = [p for p in apis.world.payments if p.method == "Credit Card"]
    assert payments
    assert by_payment.keys() == {p.id for p in payments}
    invoices = {i.id: i for i in apis.world.invoices}
    for payment in payments:
        invoice, charge = invoices[payment.invoice_id], by_payment[payment.id]
        assert charge["metadata"]["invoice_number"] == invoice.number
        assert charge["amount"] == round(payment.amount * 100)
        assert charge["amount_captured"] == charge["amount"]
        assert charge["paid"] and charge["status"] == "succeeded"
        assert charge["currency"] == apis.world.org.currency.lower()
        assert client_ids[charge["customer"]] == invoice.client_id == payment.client_id
        if invoice.status == "Paid":
            assert round(invoice.total * 100) == charge["amount"]


def text_words(text: str) -> str:
    return " ".join(text.split())


def extracted(content: bytes, kind: str) -> str:
    """Real parsers catch empty blobs and format-specific content loss."""
    stream = io.BytesIO(content)
    if kind == "docx":
        doc = DocxDocument(stream)
        return text_words("\n".join(p.text for p in doc.paragraphs))
    if kind == "pdf":
        return text_words("\n".join(p.extract_text() for p in PdfReader(stream).pages))
    if kind == "xlsx":
        workbook = load_workbook(stream, read_only=True, data_only=True)
        try:
            return text_words(
                "\n".join(
                    "\t".join(cell_text(value) for value in row)
                    for sheet in workbook
                    for row in sheet.iter_rows(values_only=True)
                )
            )
        finally:
            workbook.close()
    raise AssertionError(kind)


def cell_text(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time().isoformat() == "00:00:00" else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value) if value is not None else ""


@dataclass
class StoredFile:
    provider: str
    metadata_url: str
    download_url: str
    auth: dict[str, str]
    metadata_method: str = "GET"
    metadata_body: dict[str, Any] | None = None
    metadata_params: dict[str, Any] | None = None
    download_method: str = "GET"
    download_params: dict[str, Any] | None = None
    download_headers: dict[str, str] | None = None

    async def download(self, apis: APIs) -> bytes:
        response = await apis.client.request(
            self.download_method,
            self.download_url,
            headers={**self.auth, **(self.download_headers or {})},
            params=self.download_params,
            follow_redirects=True,
        )
        assert response.is_success, (self.provider, response.status_code, response.text[:300])
        assert response.content
        return response.content


async def storage_files(apis: APIs, doc: Document) -> list[StoredFile]:
    """Discover vendor IDs by names and folder paths, as a consumer must."""
    space = apis.world.space(doc.space_id)
    owner = doc.owner_id
    result: list[StoredFile] = []
    headers = await apis.auth("https://graph.microsoft.com", user=owner)
    sites = await apis.pages(
        GRAPH + "/sites", "value", headers=headers, params={"search": space.name}, next_key="@odata.nextLink"
    )
    site = next(s for s in sites if s["displayName"] == space.name)
    drive = await apis.get(GRAPH + f"/sites/{site['id']}/drive", headers=headers)
    item_url = f"{GRAPH}/drives/{drive['id']}/root:/{doc.path}"
    item = await apis.get(item_url, headers=headers)
    url = f"{GRAPH}/drives/{drive['id']}/items/{item['id']}"
    result.append(StoredFile("microsoft", url, url + "/content", headers))

    headers = await apis.auth(GOOGLE, user=owner)
    drives = await apis.pages(
        GOOGLE + "/drive/v3/drives", "drives", headers=headers, params={"pageSize": 100}
    )
    drive = next(d for d in drives if d["name"] == space.name)
    parent = drive["id"]
    names = [*(doc.folder.split("/") if doc.folder else []), doc.name]
    for index, name in enumerate(names):
        escaped = name.replace("\\", "\\\\").replace("'", "\\'")
        name_query = f"name = '{escaped}'"
        if index == len(names) - 1:
            stem = escaped.rsplit(".", 1)[0]
            name_query = f"({name_query} or name = '{stem}')"
        files = await apis.pages(
            GOOGLE + "/drive/v3/files",
            "files",
            headers=headers,
            params={
                "q": f"{name_query} and '{parent}' in parents and trashed = false",
                "fields": "nextPageToken,files(id,name,mimeType)",
                "pageSize": 100,
                "includeItemsFromAllDrives": "true",
                "supportsAllDrives": "true",
            },
        )
        assert len(files) == 1, (doc.id, name, files)
        parent = files[0]["id"]
    url = GOOGLE + "/drive/v3/files/" + parent
    native = files[0]["mimeType"].startswith("application/vnd.google-apps.")
    result.append(
        StoredFile(
            "google",
            url,
            url + "/export" if native else url,
            headers,
            metadata_params={"supportsAllDrives": "true"},
            download_params={"mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}
            if native
            else {"alt": "media", "supportsAllDrives": "true"},
        )
    )

    headers = await apis.auth(BOX, user=owner)
    page = await apis.get(
        BOX + "/2.0/search",
        headers=headers,
        params={
            "query": '"' + doc.name + '"',
            "type": "file",
            "content_types": "name",
            "fields": "id,name,path_collection",
            "limit": 200,
        },
    )
    suffix = [space.name, *(doc.folder.split("/") if doc.folder else [])]
    matches = [
        f
        for f in page["entries"]
        if f["name"] == doc.name
        and [e["name"] for e in f["path_collection"]["entries"]][-len(suffix) :] == suffix
    ]
    assert len(matches) == 1, (doc.id, suffix, page)
    url = BOX + "/2.0/files/" + matches[0]["id"]
    result.append(StoredFile("box", url, url + "/content", headers))

    headers = await apis.auth(DROPBOX, user=owner)
    account = await apis.post(DROPBOX + "/2/users/get_current_account", headers=headers)
    headers["Dropbox-API-Path-Root"] = json.dumps(
        {".tag": "root", "root": account["root_info"]["root_namespace_id"]}
    )
    page = await apis.post(
        DROPBOX + "/2/files/search_v2",
        headers=headers,
        json={"query": doc.name, "options": {"max_results": 1000, "filename_only": True}},
    )
    files = [m["metadata"]["metadata"] for m in page["matches"]]
    while page["has_more"]:
        page = await apis.post(
            DROPBOX + "/2/files/search/continue_v2", headers=headers, json={"cursor": page["cursor"]}
        )
        files.extend(m["metadata"]["metadata"] for m in page["matches"])
    matches = [f for f in files if f["path_display"].endswith("/" + doc.path)]
    assert len(matches) == 1, (doc.id, files)
    file_id = matches[0]["id"]
    result.append(
        StoredFile(
            "dropbox",
            DROPBOX + "/2/files/get_metadata",
            "https://content.dropboxapi.com/2/files/download",
            headers,
            metadata_method="POST",
            metadata_body={"path": file_id},
            download_method="POST",
            download_headers={"Dropbox-API-Arg": json.dumps({"path": file_id})},
        )
    )
    return result


@pytest.mark.parametrize("kind", ["docx", "pdf", "xlsx"])
async def test_document_downloads_have_identical_content(apis: APIs, kind: str) -> None:
    doc = next(
        d for d in apis.world.documents if d.kind == kind and d.space_id == "clients" and not d.trashed
    )
    files = await storage_files(apis, doc)
    texts = {file.provider: extracted(await file.download(apis), kind) for file in files}
    assert len(set(texts.values())) == 1, texts
    text = next(iter(texts.values()))
    assert text
    if kind == "xlsx":
        assert doc.table
        assert text == text_words("\n".join("\t".join(str(c) for c in row) for row in doc.table))
    else:
        for paragraph in doc.paragraphs:
            # PDF's standard fonts encode Windows-1252; Office preserves Unicode.
            expected = (
                paragraph.encode("cp1252", errors="replace").decode("cp1252") if kind == "pdf" else paragraph
            )
            assert text_words(expected) in text


async def wiki_pages(apis: APIs, doc: Document) -> list[tuple[str, str, dict[str, str]]]:
    notion = "https://api.notion.com"
    headers = {**await apis.auth(notion, user=doc.owner_id), "Notion-Version": "2025-09-03"}
    page = await apis.post(
        notion + "/v1/search", headers=headers, json={"query": doc.title, "page_size": 100}
    )
    matches = [r for r in page["results"] if r["object"] == "page"]
    assert len(matches) == 1, (doc.id, page)
    notion_url = notion + "/v1/pages/" + matches[0]["id"]
    atlassian = "https://halcyonadvisory.atlassian.net"
    confluence_headers = await apis.auth(atlassian, user=doc.owner_id)
    page = await apis.get(
        atlassian + "/wiki/api/v2/pages",
        headers=confluence_headers,
        params={"title": doc.title, "limit": 100},
    )
    assert len(page["results"]) == 1, (doc.id, page)
    confluence_url = atlassian + "/wiki/api/v2/pages/" + page["results"][0]["id"]
    return [("notion", notion_url, headers), ("atlassian", confluence_url, confluence_headers)]


class HTMLText(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        self.parts.append(data)


async def test_documents_have_matching_notion_and_confluence_pages(apis: APIs) -> None:
    doc = next(
        d for d in apis.world.documents if d.kind == "docx" and d.space_id == "clients" and not d.trashed
    )
    pages = await wiki_pages(apis, doc)
    for provider, url, headers in pages:
        if provider == "notion":
            page_id = url.rsplit("/", 1)[1]
            blocks = await apis.pages(
                f"https://api.notion.com/v1/blocks/{page_id}/children",
                "results",
                headers=headers,
                next_key="next_cursor",
                token_param="start_cursor",
            )
            text = "\n".join(
                "".join(t["plain_text"] for t in b[b["type"]].get("rich_text", [])) for b in blocks
            )
        else:
            page = await apis.get(url, headers=headers, params={"body-format": "storage"})
            parser = HTMLText()
            parser.feed(page["body"]["storage"]["value"])
            text = "\n".join(parser.parts)
        assert text_words(doc.title) in text_words(text), provider
        for paragraph in doc.paragraphs:
            assert text_words(paragraph) in text_words(text), (provider, paragraph)


async def test_restricted_document_is_hidden_from_nonmembers(apis: APIs) -> None:
    doc = next(
        d for d in apis.world.documents if d.restricted and d.space_id == "finance" and d.kind == "xlsx"
    )
    allowed_ids = {doc.owner_id}
    for share in doc.shares:
        if share.principal_id:
            allowed_ids.update(apis.world.expand_principal(share.principal_id))
    members = {
        pid
        for principal, _ in apis.world.space(doc.space_id).members
        for pid in apis.world.expand_principal(principal)
    }
    outsider = next(
        p
        for p in apis.world.active_employees
        if p.id in members and p.id not in allowed_ids and not p.is_admin and not p.is_partner
    )
    files = await storage_files(apis, doc)
    for file in files:
        assert extracted(await file.download(apis), doc.kind)
        host = str(httpx.URL(file.metadata_url).copy_with(path="", query=None))
        denied_auth = await apis.auth(host, user=outsider.id)
        response = await apis.client.request(
            file.metadata_method,
            file.metadata_url,
            headers=denied_auth,
            json=file.metadata_body,
            params=file.metadata_params,
        )
        # Dropbox uses 409 path/not_found, while the other vendors use 403/404.
        assert response.status_code in ({409} if file.provider == "dropbox" else {403, 404}), (
            file.provider,
            response.text,
        )
        if file.provider == "dropbox":
            assert response.json()["error"]["path"][".tag"] == "not_found"
        download = await apis.client.request(
            file.download_method,
            file.download_url,
            headers={**denied_auth, **(file.download_headers or {})},
            params=file.download_params,
            follow_redirects=True,
        )
        assert download.status_code in ({409} if file.provider == "dropbox" else {403, 404}), (
            file.provider,
            download.text[:300],
        )
    for provider, url, headers in await wiki_pages(apis, doc):
        assert (await apis.client.get(url, headers=headers)).is_success
        host = str(httpx.URL(url).copy_with(path="", query=None))
        denied = {**await apis.auth(host, user=outsider.id)}
        if provider == "notion":
            denied["Notion-Version"] = headers["Notion-Version"]
        response = await apis.client.get(url, headers=denied)
        assert response.status_code in (403, 404), (provider, response.text)


def instant(value: str) -> datetime:
    parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
    return parsed.replace(tzinfo=UTC) if parsed.tzinfo is None else parsed.astimezone(UTC)


async def calendar_events(apis: APIs, meeting: Meeting) -> tuple[dict[str, Any], dict[str, Any]]:
    start, end = meeting.start - timedelta(minutes=1), meeting.end + timedelta(minutes=1)
    graph = await apis.pages(
        GRAPH + "/me/calendarView",
        "value",
        headers=await apis.auth("https://graph.microsoft.com", user=meeting.organizer_id),
        params={"startDateTime": start.isoformat(), "endDateTime": end.isoformat(), "$top": 100},
        next_key="@odata.nextLink",
    )
    outlook = next(
        e for e in graph if e["subject"] == meeting.title and instant(e["start"]["dateTime"]) == meeting.start
    )
    google = await apis.pages(
        GOOGLE + "/calendar/v3/calendars/primary/events",
        "items",
        headers=await apis.auth(GOOGLE, user=meeting.organizer_id),
        params={
            "singleEvents": "true",
            "timeMin": start.isoformat(),
            "timeMax": end.isoformat(),
            "maxResults": 100,
        },
    )
    event = next(
        e
        for e in google
        if e["summary"] == meeting.title and instant(e["start"]["dateTime"]) == meeting.start
    )
    organizer = apis.world.person(meeting.organizer_id).email
    assert outlook["organizer"]["emailAddress"]["address"] == event["organizer"]["email"] == organizer
    assert instant(outlook["end"]["dateTime"]) == instant(event["end"]["dateTime"]) == meeting.end
    expected = {apis.world.person(p).email for p in (*meeting.attendee_ids, *meeting.optional_ids)}
    assert {a["emailAddress"]["address"] for a in outlook["attendees"]} == expected
    assert {a["email"] for a in event["attendees"]} - {organizer} == expected
    return outlook, event


def transcript_turns(apis: APIs, meeting: Meeting) -> list[tuple[str, str]]:
    return [(apis.world.person(t.speaker_id).display_name, text_words(t.text)) for t in meeting.transcript]


def logical_turns(turns: list[tuple[str, str]]) -> list[tuple[str, str]]:
    # VTT cue duration limits split a speaker's utterance differently by platform.
    return [
        (speaker, " ".join(text for _, text in group))
        for speaker, group in groupby(turns, key=lambda t: t[0])
    ]


async def teams_transcript(apis: APIs, meeting: Meeting, outlook: dict[str, Any]) -> list[tuple[str, str]]:
    headers = await apis.auth("https://graph.microsoft.com", user=meeting.organizer_id)
    link = outlook["onlineMeeting"]["joinUrl"]
    rows = await apis.pages(
        GRAPH + "/me/onlineMeetings",
        "value",
        headers=headers,
        params={"$filter": "JoinWebUrl eq '" + link.replace("'", "''") + "'"},
        next_key="@odata.nextLink",
    )
    assert len(rows) == 1 and rows[0]["subject"] == meeting.title
    url = GRAPH + "/me/onlineMeetings/" + rows[0]["id"]
    transcripts = await apis.get(url + "/transcripts", headers=headers)
    assert len(transcripts["value"]) == 1
    response = await apis.client.get(
        url + "/transcripts/" + transcripts["value"][0]["id"] + "/content", headers=headers
    )
    response.raise_for_status()
    return [(cue.speaker, text_words(cue.text)) for cue in media.parse_vtt(response.content)]


async def zoom_transcript(
    apis: APIs, meeting: Meeting, outlook: dict[str, Any], google: dict[str, Any]
) -> list[tuple[str, str]]:
    headers = await apis.auth(ZOOM, user=meeting.organizer_id)
    recordings = await apis.pages(
        ZOOM + "/v2/users/me/recordings",
        "meetings",
        headers=headers,
        params={
            "from": meeting.start.date().isoformat(),
            "to": (meeting.start.date() + timedelta(days=1)).isoformat(),
            "page_size": 100,
        },
        next_key="next_page_token",
        token_param="next_page_token",
    )
    recording = next(r for r in recordings if r["topic"] == meeting.title)
    detail = await apis.get(ZOOM + f"/v2/meetings/{recording['id']}", headers=headers)
    # A calendar link must join the very meeting whose artifacts we download.
    assert google["location"] == detail["join_url"]
    assert detail["join_url"] in outlook["body"]["content"]
    transcript = next(r for r in recording["recording_files"] if r["file_type"] == "TRANSCRIPT")
    response = await apis.client.get(transcript["download_url"], headers=headers, follow_redirects=True)
    response.raise_for_status()
    uuid = recording["uuid"]
    from urllib.parse import quote

    encoded = (
        quote(quote(uuid, safe=""), safe="") if uuid.startswith("/") or "//" in uuid else quote(uuid, safe="")
    )
    participants = await apis.pages(
        ZOOM + f"/v2/past_meetings/{encoded}/participants",
        "participants",
        headers=headers,
        params={"page_size": 100},
        next_key="next_page_token",
        token_param="next_page_token",
    )
    assert participants
    contacts = await apis.salesforce("SELECT Email, Name FROM Contact")
    directory = await apis.pages(
        "https://admin.googleapis.com/admin/directory/v1/users",
        "users",
        headers=await apis.auth(GOOGLE, user="p035"),
        params={"customer": "my_customer", "maxResults": 500},
    )
    emails = {p["primaryEmail"] for p in directory} | {p["Email"] for p in contacts}
    names = {p["name"]["fullName"] for p in directory} | {p["Name"] for p in contacts}
    for participant in participants:
        assert participant["name"] in names
        if participant.get("user_email"):
            assert participant["user_email"] in emails
    return [(cue.speaker, text_words(cue.text)) for cue in media.parse_vtt(response.content)]


async def meet_transcript(
    apis: APIs, meeting: Meeting, outlook: dict[str, Any], google: dict[str, Any]
) -> list[tuple[str, str]]:
    host = "https://meet.googleapis.com"
    headers = await apis.auth(host, user=meeting.organizer_id)
    assert google["hangoutLink"] in outlook["body"]["content"]
    space = await apis.get(
        host + "/v2/spaces/" + urlsplit(google["hangoutLink"]).path.strip("/"), headers=headers
    )
    conferences = await apis.pages(
        host + "/v2/conferenceRecords",
        "conferenceRecords",
        headers=headers,
        params={
            "filter": f'space.name="{space["name"]}" AND start_time>="{meeting.start.isoformat()}"',
            "pageSize": 100,
        },
    )
    conference = next(c for c in conferences if instant(c["startTime"]) == meeting.start)
    url = host + "/v2/" + conference["name"]
    transcripts = await apis.get(url + "/transcripts", headers=headers)
    transcript = transcripts["transcripts"][0]
    entries = await apis.pages(
        host + "/v2/" + transcript["name"] + "/entries",
        "transcriptEntries",
        headers=headers,
        params={"pageSize": 100},
    )
    participants = await apis.pages(
        url + "/participants", "participants", headers=headers, params={"pageSize": 100}
    )
    people = {p["name"]: p for p in participants}
    directory = await apis.pages(
        "https://admin.googleapis.com/admin/directory/v1/users",
        "users",
        headers=await apis.auth(GOOGLE, user="p035"),
        params={"customer": "my_customer", "maxResults": 500},
    )
    users = {"users/" + p["id"]: p for p in directory}
    contacts = await apis.salesforce("SELECT Email, Name FROM Contact")
    known_names = {c["Name"] for c in contacts} | {p["name"]["fullName"] for p in directory}
    for participant in participants:
        identity = (
            participant.get("signedinUser") or participant.get("anonymousUser") or participant["phoneUser"]
        )
        assert identity["displayName"] in known_names
        if "user" in identity:
            assert identity["user"] in users
            assert users[identity["user"]]["name"]["fullName"] == identity["displayName"]
    turns = []
    for entry in entries:
        p = people[entry["participant"]]
        identity = p.get("signedinUser") or p.get("anonymousUser") or p["phoneUser"]
        turns.append((identity["displayName"], text_words(entry["text"])))
    doc_id = transcript["docsDestination"]["document"]
    response = await apis.client.get(
        GOOGLE + f"/drive/v3/files/{doc_id}/export",
        headers=headers,
        params={"mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
    )
    response.raise_for_status()
    text = extracted(response.content, "docx")
    for speaker, utterance in turns:
        assert f"{speaker}: {utterance}" in text
    return turns


@pytest.mark.parametrize("platform", ["teams", "zoom", "meet"])
async def test_recorded_meeting_transcript_and_both_calendars(apis: APIs, platform: str) -> None:
    meeting = next(
        m
        for m in apis.world.meetings
        if m.platform == platform
        and m.recorded
        and m.transcribed
        and m.kind == "client"
        and not m.cancelled
        and m.start >= apis.world.anchor - timedelta(days=20)
    )
    outlook, google = await calendar_events(apis, meeting)
    if platform == "teams":
        assert outlook["onlineMeeting"]["joinUrl"] in google["description"]
        turns = await teams_transcript(apis, meeting, outlook)
    elif platform == "zoom":
        turns = await zoom_transcript(apis, meeting, outlook, google)
    else:
        turns = await meet_transcript(apis, meeting, outlook, google)
    assert logical_turns(turns) == logical_turns(transcript_turns(apis, meeting))


async def test_webex_projection_has_the_original_meeting_transcript(apis: APIs) -> None:
    clients = {
        c.id
        for c in apis.world.clients
        if c.name in {"Northwind Logistics", "Atlas Freight Partners", "Meridian Legal"}
    }
    meeting = next(
        m
        for m in apis.world.meetings
        if m.client_id in clients
        and m.happened
        and m.recorded
        and m.transcribed
        and not m.cancelled
        and m.start >= apis.world.anchor - timedelta(days=20)
    )
    outlook, google = await calendar_events(apis, meeting)
    headers = await apis.auth(WEBEX, user=meeting.organizer_id)
    rows = await apis.pages(
        WEBEX + "/v1/meetings",
        "items",
        headers=headers,
        params={
            "meetingType": "meeting",
            "from": (meeting.start - timedelta(minutes=1)).isoformat(),
            "to": (meeting.end + timedelta(minutes=1)).isoformat(),
            "max": 100,
        },
    )
    row = next(r for r in rows if r["title"] == meeting.title and instant(r["start"]) == meeting.start)
    transcripts = await apis.get(
        WEBEX + "/v1/meetingTranscripts", headers=headers, params={"meetingId": row["id"]}
    )
    transcript = transcripts["items"][0]
    response = await apis.client.get(
        WEBEX + f"/v1/meetingTranscripts/{transcript['id']}/download", headers=headers
    )
    response.raise_for_status()
    turns = [(cue.speaker, text_words(cue.text)) for cue in media.parse_vtt(response.content)]
    assert logical_turns(turns) == logical_turns(transcript_turns(apis, meeting))
    if meeting.platform == "teams":
        assert logical_turns(turns) == logical_turns(await teams_transcript(apis, meeting, outlook))
    elif meeting.platform == "zoom":
        assert logical_turns(turns) == logical_turns(await zoom_transcript(apis, meeting, outlook, google))
    else:
        assert logical_turns(turns) == logical_turns(await meet_transcript(apis, meeting, outlook, google))


async def contact_numbers(apis: APIs) -> dict[str, str]:
    salesforce = {
        r["Email"]: r for r in await apis.salesforce("SELECT Email, Name, Phone, MobilePhone FROM Contact")
    }
    hubspot = {
        r["properties"]["email"]: r["properties"]
        for r in await apis.hubspot("contacts", "email,firstname,lastname,phone,mobilephone")
    }
    numbers: dict[str, str] = {}
    for client in apis.world.clients:
        for pid in client.contact_ids:
            person = apis.world.person(pid)
            sf, hs = salesforce[person.email], hubspot[person.email]
            assert sf["Name"] == f"{hs['firstname']} {hs['lastname']}" == person.display_name
            assert sf["MobilePhone"] == hs.get("mobilephone") == person.phone_mobile
            assert sf["Phone"] == hs.get("phone") == person.phone_business
            if person.phone_mobile:
                numbers[person.phone_mobile] = person.email
    assert numbers
    return numbers


async def test_sms_and_whatsapp_history_resolve_to_crm_contacts(apis: APIs) -> None:
    numbers = await contact_numbers(apis)
    host = "https://api.twilio.com"
    credentials = await apis.post(host + "/_emulator/tokens", json={})
    url = host + f"/2010-04-01/Accounts/{credentials['account_sid']}/Messages.json"
    auth = httpx.BasicAuth(credentials["account_sid"], credentials["auth_token"])
    messages: list[dict[str, Any]] = []
    params: dict[str, Any] | None = {"PageSize": 31}
    while True:
        page = await apis.get(url, auth=auth, params=params)
        messages.extend(page["messages"])
        if page["next_page_uri"] is None:
            break
        url, params = host + page["next_page_uri"], None
    assert messages
    for message in messages:
        contact = message["from"] if message["direction"] == "inbound" else message["to"]
        assert contact.removeprefix("whatsapp:") in numbers, message["sid"]
    # Every world thread is represented, including both directions and media messages.
    for conversation in apis.world.conversations:
        if conversation.channel not in ("sms", "whatsapp"):
            continue
        prefix = "whatsapp:" if conversation.channel == "whatsapp" else ""
        person = apis.world.person(conversation.person_id)
        for source in conversation.messages:
            outbound = source.direction == "outbound"
            expected_from = prefix + (conversation.business_number if outbound else person.phone_mobile)
            expected_to = prefix + (person.phone_mobile if outbound else conversation.business_number)
            matches = [
                m
                for m in messages
                if m["body"] == source.text and m["from"] == expected_from and m["to"] == expected_to
            ]
            assert matches, (conversation.id, source.id)


async def test_meta_whatsapp_webhook_resolves_to_the_same_crm_contact(apis: APIs) -> None:
    # The Cloud API has no public history endpoint. Consume its documented inbound
    # webhook while Twilio's public API above verifies seeded WhatsApp history.
    numbers = await contact_numbers(apis)
    received: list[dict[str, Any]] = []
    secret = "synthetic-cross-system-webhook-secret"
    receiver = FastAPI()

    @receiver.get("/whatsapp")
    async def challenge(request: Request) -> PlainTextResponse:
        assert request.query_params["hub.verify_token"] == "cross-system-verifier"
        return PlainTextResponse(request.query_params["hub.challenge"])

    @receiver.post("/whatsapp")
    async def webhook(request: Request) -> dict[str, bool]:
        body = await request.body()
        signature = "sha256=" + hmac.new(secret.encode(), body, hashlib.sha256).hexdigest()
        assert hmac.compare_digest(request.headers["x-hub-signature-256"], signature)
        received.append(json.loads(body))
        return {"received": True}

    apis.client.hub.webhook_transport = httpx.ASGITransport(app=receiver)
    host = "https://graph.facebook.com"
    configured = await apis.post(
        host + "/_emulator/actions/configure_webhook",
        json={
            "object": "whatsapp_business_account",
            "callback_url": "https://consumer.example/whatsapp",
            "verify_token": "cross-system-verifier",
            "app_secret": secret,
            "fields": ["messages"],
        },
    )
    assert configured["success"]
    conversations = [c for c in apis.world.conversations if c.channel == "whatsapp"]
    assert conversations
    for conversation in conversations:
        person = apis.world.person(conversation.person_id)
        source = next(m for m in conversation.messages if m.direction == "inbound")
        result = await apis.post(
            host + "/_emulator/actions/inbound_whatsapp", json={"from": person.id, "text": source.text}
        )
        assert result["success"] and result["webhook_delivered"]
        value = received[-1]["entry"][0]["changes"][0]["value"]
        contact, message = value["contacts"][0], value["messages"][0]
        assert contact["wa_id"] == message["from"] == person.phone_mobile.removeprefix("+")
        assert numbers["+" + message["from"]] == person.email
        assert contact["profile"]["name"] == person.display_name
        assert message["text"]["body"] == source.text
    assert len(received) == len(conversations)
