"""Seed documents must open in the real parsers, not just look like the right file type."""

from __future__ import annotations

import csv
import io
import re
import struct
import zipfile
import zlib
from datetime import UTC, date, datetime
from xml.etree import ElementTree

import docx
import openpyxl
import pptx
import pypdf
import pytest

from emulators.world import documents as D

CREATED = datetime(2025, 3, 14, 9, 26, 52, tzinfo=UTC)
SPECIAL = 'R&D <Q3> "plan" — café'
PARAGRAPHS = [
    "First paragraph about the quarterly plan.",
    f"Second paragraph with {SPECIAL}.",
    "  leading and trailing spaces are preserved  ",
    "Fourth paragraph, ends the document.",
]
ALL_KINDS = ["docx", "xlsx", "pptx", "pdf", "csv", "md", "txt", "png"]


def _squash(text: str) -> str:
    """Collapse whitespace: PDF extraction splits wrapped paragraphs across lines."""
    return " ".join(text.split())


def _naive(stamp: datetime) -> datetime:
    return stamp.replace(tzinfo=None)


def _open_docx(data: bytes) -> docx.document.Document:
    return docx.Document(io.BytesIO(data))


def _pdf_text(reader: pypdf.PdfReader) -> str:
    return _squash("\n".join(page.extract_text() for page in reader.pages))


def _all_slide_text(slide: pptx.slide.Slide) -> list[str]:
    return [p.text for shape in slide.shapes if shape.has_text_frame for p in shape.text_frame.paragraphs]


# --------------------------------------------------------------------------------------------------
# constants
# --------------------------------------------------------------------------------------------------


def test_mime_constants_and_extension_map() -> None:
    assert D.MIME_CSV == "text/csv"
    assert D.MIME_MARKDOWN == "text/markdown"
    assert D.MIME_TEXT == "text/plain"
    assert D.MIME_PDF == "application/pdf"
    assert D.MIME_PNG == "image/png"
    assert D.MIME_DOCX.endswith("wordprocessingml.document")
    assert D.MIME_XLSX.endswith("spreadsheetml.sheet")
    assert D.MIME_PPTX.endswith("presentationml.presentation")
    expected = {
        "docx": D.MIME_DOCX,
        "xlsx": D.MIME_XLSX,
        "pptx": D.MIME_PPTX,
        "pdf": D.MIME_PDF,
        "csv": "text/csv",
        "md": "text/markdown",
        "txt": "text/plain",
        "png": "image/png",
        "json": "application/json",
        "html": "text/html",
        "jpg": "image/jpeg",
        "jpeg": "image/jpeg",
        "gif": "image/gif",
        "mp4": "video/mp4",
        "m4a": "audio/mp4",
        "wav": "audio/wav",
        "vtt": "text/vtt",
        "zip": "application/zip",
        "eml": "message/rfc822",
        "msg": "application/vnd.ms-outlook",
        "doc": "application/msword",
        "xls": "application/vnd.ms-excel",
        "ppt": "application/vnd.ms-powerpoint",
    }
    for ext, mime in expected.items():
        assert D.EXTENSION_TO_MIME[ext] == mime
    assert all(ext == ext.lower() and not ext.startswith(".") for ext in D.EXTENSION_TO_MIME)


# --------------------------------------------------------------------------------------------------
# docx
# --------------------------------------------------------------------------------------------------


def test_docx_text_and_styles_round_trip() -> None:
    data = D.make_docx("Quarterly Plan", PARAGRAPHS, author="Ada Lovelace", created=CREATED)
    doc = _open_docx(data)
    assert [p.text for p in doc.paragraphs] == ["Quarterly Plan", *PARAGRAPHS]
    assert doc.paragraphs[0].style.name == "Title"
    assert all(p.style.name == "Normal" for p in doc.paragraphs[1:])
    style_ids = {s.style_id for s in doc.styles}
    assert {"Title", "Heading1", "Normal"} <= style_ids


def test_docx_headings_are_inserted_before_their_paragraph() -> None:
    data = D.make_docx("T", ["alpha", "beta", "gamma"], headings=["Intro", "", "Wrap-up"])
    doc = _open_docx(data)
    assert [(p.style.name, p.text) for p in doc.paragraphs] == [
        ("Title", "T"),
        ("Heading 1", "Intro"),
        ("Normal", "alpha"),
        ("Normal", "beta"),
        ("Heading 1", "Wrap-up"),
        ("Normal", "gamma"),
    ]


def test_docx_core_properties() -> None:
    doc = _open_docx(D.make_docx(SPECIAL, ["x"], author="Ada & Bob", created=CREATED))
    props = doc.core_properties
    assert props.title == SPECIAL
    assert props.author == "Ada & Bob"
    assert _naive(props.created) == _naive(CREATED)
    assert _naive(props.modified) == _naive(CREATED)


def test_docx_escapes_xml_special_characters() -> None:
    doc = _open_docx(D.make_docx(SPECIAL, [SPECIAL, f"<b>{SPECIAL}</b> & ]]> end"]))
    assert doc.paragraphs[0].text == SPECIAL
    assert doc.paragraphs[1].text == SPECIAL
    assert doc.paragraphs[2].text == f"<b>{SPECIAL}</b> & ]]> end"


def test_docx_preserves_line_breaks_tabs_and_drops_illegal_control_chars() -> None:
    doc = _open_docx(D.make_docx("T", ["line one\nline two\tafter tab", "bad\x00char\x0b!"]))
    assert doc.paragraphs[1].text == "line one\nline two\tafter tab"
    assert doc.paragraphs[2].text == "badchar!"


def test_docx_package_structure_is_well_formed() -> None:
    data = D.make_docx("T", ["p"], author="a", created=CREATED)
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        names = zf.namelist()
        assert names[0] == "[Content_Types].xml"
        assert {
            "_rels/.rels",
            "word/document.xml",
            "word/styles.xml",
            "docProps/core.xml",
            "docProps/app.xml",
        } <= set(names)
        for name in names:
            ElementTree.fromstring(zf.read(name))  # every part parses as XML
        assert b'xml:space="preserve"' in zf.read("word/document.xml")
        core = zf.read("docProps/core.xml").decode()
        assert "<dc:title>T</dc:title>" in core
        assert "<dc:creator>a</dc:creator>" in core
        assert 'xsi:type="dcterms:W3CDTF">2025-03-14T09:26:52Z<' in core


# --------------------------------------------------------------------------------------------------
# xlsx
# --------------------------------------------------------------------------------------------------


def test_xlsx_values_round_trip() -> None:
    when = datetime(2025, 1, 31, 13, 45, 30)
    rows: list[list[object]] = [
        ["Name", "Qty", "Price", "Active", "Due", "At"],
        ["Widget", 3, 1.5, True, date(2025, 1, 31), when],
        [SPECIAL, -7, 0.1, False, None, datetime(2024, 2, 29, 0, 0, 0)],
        ["  padded  ", 10**12, 1e-7, True, date(1999, 12, 31), when],
    ]
    wb = openpyxl.load_workbook(io.BytesIO(D.make_xlsx({"Data": rows}, author="Ada", created=CREATED)))
    ws = wb["Data"]
    assert ws["A1"].value == "Name"
    assert ws["A2"].value == "Widget"
    assert ws["A3"].value == SPECIAL
    assert ws["A4"].value == "  padded  "
    assert ws["B2"].value == 3 and isinstance(ws["B2"].value, int)
    assert ws["B3"].value == -7
    assert ws["B4"].value == 10**12
    assert ws["C2"].value == 1.5
    assert ws["C3"].value == 0.1
    assert ws["C4"].value == 1e-7
    assert ws["D2"].value is True
    assert ws["D3"].value is False
    assert ws["E2"].value == datetime(2025, 1, 31)
    assert ws["E2"].is_date
    assert ws["E3"].value is None
    assert ws["E4"].value == datetime(1999, 12, 31)
    assert ws["F2"].value == when
    assert ws["F2"].is_date
    assert ws["F3"].value == datetime(2024, 2, 29)
    assert ws.max_row == 4 and ws.max_column == 6


def test_xlsx_header_row_is_bold_and_data_rows_are_not() -> None:
    ws = openpyxl.load_workbook(io.BytesIO(D.make_xlsx({"S": [["h1", "h2"], ["a", 1]]}))).active
    assert ws["A1"].font.b and ws["B1"].font.b
    assert not ws["A2"].font.b and not ws["B2"].font.b


def test_xlsx_timezone_aware_datetimes_are_stored_as_utc() -> None:
    from datetime import timedelta, timezone

    ist = timezone(timedelta(hours=5, minutes=30))
    ws = openpyxl.load_workbook(
        io.BytesIO(D.make_xlsx({"S": [["when"], [datetime(2025, 6, 1, 12, 0, 0, tzinfo=ist)]]}))
    ).active
    assert ws["A2"].value == datetime(2025, 6, 1, 6, 30, 0)


def test_xlsx_workbook_properties_and_multiple_sheets() -> None:
    data = D.make_xlsx({"First": [["a"], [1]], "Second": [["b"], ["x"]]}, author="Grace", created=CREATED)
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["First", "Second"]
    assert wb["Second"]["A2"].value == "x"
    assert wb.properties.creator == "Grace"
    assert _naive(wb.properties.created) == _naive(CREATED)


def test_xlsx_sheet_names_are_sanitised_to_excel_rules() -> None:
    long = "A very long sheet name that exceeds thirty-one characters"
    raw = {
        "Q1/Q2: [Plan]*?\\ok": [["x"]],
        long: [["x"]],
        long.upper(): [["x"]],  # collides with the previous one once truncated, ignoring case
        "": [["x"]],
        "'quoted'": [["x"]],
    }
    wb = openpyxl.load_workbook(io.BytesIO(D.make_xlsx(raw)))
    assert len(wb.sheetnames) == 5
    assert len({n.lower() for n in wb.sheetnames}) == 5
    for name in wb.sheetnames:
        assert 1 <= len(name) <= 31
        assert not set(name) & set("[]:*?/\\")
        assert not name.startswith("'") and not name.endswith("'")
    assert wb.sheetnames[0] == "Q1_Q2_ _Plan____ok"
    assert wb.sheetnames[3] == "Sheet4"  # empty name falls back to its position
    assert wb.sheetnames[4] == "quoted"


def test_xlsx_empty_workbook_still_has_one_sheet() -> None:
    wb = openpyxl.load_workbook(io.BytesIO(D.make_xlsx({})))
    assert wb.sheetnames == ["Sheet1"]
    wb = openpyxl.load_workbook(io.BytesIO(D.make_xlsx({"Empty": []})))
    assert wb["Empty"].max_row == 1 and wb["Empty"]["A1"].value is None


def test_xlsx_non_finite_floats_and_ragged_rows() -> None:
    ws = openpyxl.load_workbook(
        io.BytesIO(D.make_xlsx({"S": [["a", "b", "c"], [float("nan")], [1, None, 3]]}))
    ).active
    assert ws["A2"].value == "#NUM!"
    assert ws["A3"].value == 1 and ws["B3"].value is None and ws["C3"].value == 3


def test_xlsx_package_structure_is_well_formed() -> None:
    data = D.make_xlsx({"S": [["a", 1, date(2025, 1, 1)]]}, author="x", created=CREATED)
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        names = zf.namelist()
        assert names[0] == "[Content_Types].xml"
        assert {"xl/workbook.xml", "xl/styles.xml", "xl/worksheets/sheet1.xml", "docProps/core.xml"} <= set(
            names
        )
        for name in names:
            ElementTree.fromstring(zf.read(name))
        assert b"<numFmt " in zf.read("xl/styles.xml")


# --------------------------------------------------------------------------------------------------
# pptx
# --------------------------------------------------------------------------------------------------


def test_pptx_every_slide_text_is_extractable() -> None:
    slides = [
        ("Overview", ["alpha point", "beta point"]),
        ("Details", [SPECIAL, "  spaced  ", "gamma"]),
        ("Empty", []),
    ]
    prs = pptx.Presentation(io.BytesIO(D.make_pptx("Deck Title", slides, author="Ada", created=CREATED)))
    assert len(prs.slides) == 4
    title_slide = prs.slides[0]
    assert title_slide.shapes.title.text == "Deck Title"
    assert title_slide.slide_layout.name == "Title Slide"
    assert _all_slide_text(title_slide) == ["Deck Title", "Ada"]
    for slide, (slide_title, bullets) in zip(list(prs.slides)[1:], slides, strict=True):
        assert slide.slide_layout.name == "Title and Content"
        assert slide.shapes.title.text == slide_title
        body = slide.placeholders[1].text_frame
        assert [p.text for p in body.paragraphs] == (bullets or [""])  # never a paragraph-less text body


def test_pptx_core_properties_and_size() -> None:
    prs = pptx.Presentation(
        io.BytesIO(D.make_pptx(SPECIAL, [("S", ["b"])], author="Grace & Co", created=CREATED))
    )
    props = prs.core_properties
    assert props.title == SPECIAL
    assert props.author == "Grace & Co"
    assert _naive(props.created) == _naive(CREATED)
    assert (prs.slide_width, prs.slide_height) == (12192000, 6858000)
    assert len(prs.slide_masters) == 1 and len(prs.slide_layouts) == 2


def test_pptx_without_author_has_no_subtitle_and_special_chars_survive() -> None:
    prs = pptx.Presentation(io.BytesIO(D.make_pptx(SPECIAL, [(SPECIAL, [SPECIAL, "a < b & c"])])))
    assert _all_slide_text(prs.slides[0]) == [SPECIAL]
    assert _all_slide_text(prs.slides[1]) == [SPECIAL, SPECIAL, "a < b & c"]


def test_pptx_package_structure_is_well_formed() -> None:
    data = D.make_pptx("T", [("S", ["b"])], created=CREATED)
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        names = zf.namelist()
        assert names[0] == "[Content_Types].xml"
        for required in (
            "ppt/presentation.xml",
            "ppt/slideMasters/slideMaster1.xml",
            "ppt/slideLayouts/slideLayout1.xml",
            "ppt/theme/theme1.xml",
            "ppt/slides/slide1.xml",
            "ppt/slides/slide2.xml",
            "ppt/slides/_rels/slide2.xml.rels",
        ):
            assert required in names
        for name in names:
            ElementTree.fromstring(zf.read(name))
        # every relationship target resolves to a part in the package
        for name in names:
            if not name.endswith(".rels"):
                continue
            base = name.replace("_rels/", "").removesuffix(".rels")
            base_dir = base.rpartition("/")[0]
            for rel in ElementTree.fromstring(zf.read(name)):
                parts: list[str] = [p for p in base_dir.split("/") if p]
                for seg in rel.attrib["Target"].split("/"):
                    if seg == "..":
                        parts.pop()
                    elif seg:
                        parts.append(seg)
                assert "/".join(parts) in names, f"{name} -> {rel.attrib['Target']}"


# --------------------------------------------------------------------------------------------------
# pdf
# --------------------------------------------------------------------------------------------------


def test_pdf_text_and_metadata_round_trip() -> None:
    data = D.make_pdf("Quarterly Plan", PARAGRAPHS, author="Ada Lovelace", created=CREATED)
    assert data.startswith(b"%PDF-1.4")
    assert data.rstrip().endswith(b"%%EOF")
    reader = pypdf.PdfReader(io.BytesIO(data), strict=True)
    assert len(reader.pages) == 1
    box = reader.pages[0].mediabox
    assert (float(box.width), float(box.height)) == (612.0, 792.0)
    text = _pdf_text(reader)
    assert "Quarterly Plan" in text
    for paragraph in PARAGRAPHS:
        assert _squash(paragraph) in text
    meta = reader.metadata
    assert meta is not None
    assert meta.title == "Quarterly Plan"
    assert meta.author == "Ada Lovelace"
    assert meta.producer == "emulators"
    assert meta.creation_date == CREATED
    assert b"/CreationDate (D:20250314092652Z)" in data


def test_pdf_special_characters_and_escaping() -> None:
    paragraphs = [
        SPECIAL,
        r"call f(x) = (a \ b) and ((nested) parens) \) done",
        "café naïve üñ • €5 “smart”",
    ]
    reader = pypdf.PdfReader(io.BytesIO(D.make_pdf(SPECIAL, paragraphs, author="Zoë")), strict=True)
    text = _pdf_text(reader)
    assert text.startswith(SPECIAL)  # title
    for paragraph in paragraphs:
        assert _squash(paragraph) in text
    assert reader.metadata is not None
    assert reader.metadata.title == SPECIAL  # non-ASCII metadata goes out as UTF-16BE text strings
    assert reader.metadata.author == "Zoë"


def test_pdf_characters_outside_winansi_become_question_marks() -> None:
    reader = pypdf.PdfReader(
        io.BytesIO(D.make_pdf("t", ["日本語 text \U0001f600 end", "x\x00y\x07z"])), strict=True
    )
    text = _pdf_text(reader)
    assert "??? text ? end" in text
    assert "xyz" in text


def test_pdf_long_input_spans_multiple_pages_and_wraps() -> None:
    paragraphs = [
        f"Paragraph {i:03d}: " + "lorem ipsum dolor sit amet consectetur adipiscing elit " * 4 + f"END{i:03d}"
        for i in range(60)
    ]
    reader = pypdf.PdfReader(io.BytesIO(D.make_pdf("Long report", paragraphs)), strict=True)
    assert len(reader.pages) >= 3
    text = _pdf_text(reader)
    for i in range(60):
        assert f"Paragraph {i:03d}:" in text
        assert f"END{i:03d}" in text
    # paragraphs are wrapped to ~90 characters, never emitted as one long line
    for page in reader.pages:
        lines = [ln for ln in page.extract_text().splitlines() if ln.strip()]
        assert lines
        assert max(len(ln) for ln in lines) <= 110
    first_page_lines = reader.pages[0].extract_text().splitlines()
    assert first_page_lines[0] == "Long report"


def test_pdf_unbreakable_word_is_split_instead_of_overflowing() -> None:
    word = "A" * 200
    text = _pdf_text(pypdf.PdfReader(io.BytesIO(D.make_pdf("t", [word])), strict=True))
    assert text.replace(" ", "").count("A") == 200


def test_pdf_empty_input_still_produces_one_page() -> None:
    reader = pypdf.PdfReader(io.BytesIO(D.make_pdf("", [])), strict=True)
    assert len(reader.pages) == 1


def test_pdf_xref_offsets_and_stream_lengths_are_exact() -> None:
    data = D.make_pdf("Title", [f"para {i} " * 40 for i in range(80)], author="a", created=CREATED)
    startxref = int(re.search(rb"startxref\n(\d+)\n%%EOF\n$", data).group(1))  # type: ignore[union-attr]
    assert data[startxref : startxref + 5] == b"xref\n"
    header, *entries = data[startxref:].split(b"trailer")[0].splitlines(keepends=True)[1:]
    count = int(header.split()[1])
    assert len(entries) == count
    assert all(len(entry) == 20 for entry in entries)  # fixed-width xref entries
    assert entries[0] == b"0000000000 65535 f \n"
    for number, entry in enumerate(entries[1:], start=1):
        offset = int(entry[:10])
        assert data[offset:].startswith(b"%d 0 obj\n" % number), f"object {number}"
    for match in re.finditer(rb"<< /Length (\d+) >>\nstream\n", data):
        length = int(match.group(1))
        end = match.end() + length
        assert data[end : end + 10] == b"\nendstream"
    trailer = data[data.index(b"trailer") :]
    assert b"/Size %d" % count in trailer and b"/Root 1 0 R" in trailer and b"/Info 5 0 R" in trailer


# --------------------------------------------------------------------------------------------------
# csv / markdown / text
# --------------------------------------------------------------------------------------------------


def test_csv_uses_crlf_utf8_and_quotes_when_needed() -> None:
    rows: list[list[object]] = [
        ["id", "note", "when", "flag", "price"],
        [1, 'He said "hi", then left', date(2025, 1, 31), True, 9.5],
        [2, "multi\nline", datetime(2025, 1, 31, 8, 5, 0), False, None],
        [3, "café — ünïcode", None, None, 0.1],
    ]
    data = D.make_csv(rows)
    text = data.decode("utf-8")
    assert not data.startswith(b"\xef\xbb\xbf")
    assert text.endswith("\r\n")
    assert text.split("\r\n")[0] == "id,note,when,flag,price"
    assert '"He said ""hi"", then left"' in text
    parsed = list(csv.reader(io.StringIO(text, newline="")))
    assert parsed == [
        ["id", "note", "when", "flag", "price"],
        ["1", 'He said "hi", then left', "2025-01-31", "TRUE", "9.5"],
        ["2", "multi\nline", "2025-01-31 08:05:00", "FALSE", ""],
        ["3", "café — ünïcode", "", "", "0.1"],
    ]
    # every line break outside quotes is CRLF: removing CRLFs leaves only the quoted LF
    assert text.replace("\r\n", "").count("\n") == 1
    assert D.make_csv([]) == b""


def test_markdown_structure() -> None:
    text = D.make_markdown("My Title", ["one", "two", "three"], headings=["Intro", "", "End"]).decode("utf-8")
    assert text == "# My Title\n\n## Intro\n\none\n\ntwo\n\n## End\n\nthree\n"
    assert D.make_markdown(SPECIAL, [SPECIAL]).decode("utf-8") == f"# {SPECIAL}\n\n{SPECIAL}\n"


def test_plain_text_structure() -> None:
    assert D.make_text("My Title", ["one", "two"]).decode("utf-8") == "My Title\n\none\n\ntwo\n"
    assert D.make_text(SPECIAL, [SPECIAL]).decode("utf-8") == f"{SPECIAL}\n\n{SPECIAL}\n"


# --------------------------------------------------------------------------------------------------
# png
# --------------------------------------------------------------------------------------------------


def _png_chunks(data: bytes) -> list[tuple[bytes, bytes]]:
    assert data[:8] == b"\x89PNG\r\n\x1a\n"
    chunks, pos = [], 8
    while pos < len(data):
        (length,) = struct.unpack(">I", data[pos : pos + 4])
        tag = data[pos + 4 : pos + 8]
        payload = data[pos + 8 : pos + 8 + length]
        (crc,) = struct.unpack(">I", data[pos + 8 + length : pos + 12 + length])
        assert crc == zlib.crc32(tag + payload) & 0xFFFFFFFF, f"bad CRC in {tag!r}"
        chunks.append((tag, payload))
        pos += 12 + length
    return chunks


@pytest.mark.parametrize(("width", "height"), [(64, 64), (1, 1), (13, 7), (200, 120)])
def test_png_is_valid_with_requested_dimensions(width: int, height: int) -> None:
    chunks = _png_chunks(D.make_png(width, height, seed=3))
    assert [tag for tag, _ in chunks] == [b"IHDR", b"IDAT", b"IEND"]
    assert struct.unpack(">IIBBBBB", chunks[0][1]) == (width, height, 8, 2, 0, 0, 0)
    raw = zlib.decompress(chunks[1][1])
    assert len(raw) == height * (1 + 3 * width)
    assert all(raw[row * (1 + 3 * width)] == 0 for row in range(height))  # filter byte per scanline


def test_png_defaults_and_seed_variation() -> None:
    default = D.make_png()
    assert struct.unpack(">II", default[16:24]) == (64, 64)
    assert default == D.make_png(64, 64, seed=0)
    assert D.make_png(seed=1) != D.make_png(seed=2)


@pytest.mark.parametrize(("width", "height"), [(0, 10), (10, 0), (-1, 5)])
def test_png_rejects_non_positive_dimensions(width: int, height: int) -> None:
    with pytest.raises(ValueError, match="positive"):
        D.make_png(width, height)


# --------------------------------------------------------------------------------------------------
# determinism
# --------------------------------------------------------------------------------------------------


def _build_all(created: datetime | None) -> dict[str, bytes]:
    table: list[list[object]] = [["a", "b"], [1, date(2025, 1, 1)]]
    return {
        "docx": D.make_docx("T", PARAGRAPHS, author="A", created=created, headings=["h", "", "", ""]),
        "xlsx": D.make_xlsx({"S": table}, author="A", created=created),
        "pptx": D.make_pptx("T", [("S", PARAGRAPHS)], author="A", created=created),
        "pdf": D.make_pdf("T", PARAGRAPHS, author="A", created=created),
        "csv": D.make_csv(table),
        "md": D.make_markdown("T", PARAGRAPHS, headings=["h"]),
        "txt": D.make_text("T", PARAGRAPHS),
        "png": D.make_png(32, 16, seed=9),
    }


def test_identical_inputs_give_byte_identical_output() -> None:
    assert _build_all(CREATED) == _build_all(CREATED)
    assert _build_all(None) == _build_all(None)


def test_default_created_is_fixed_and_not_the_wall_clock() -> None:
    fixed = datetime(2026, 1, 1, tzinfo=UTC)
    assert _build_all(None) == _build_all(fixed)
    # naive datetimes are interpreted as UTC
    assert _build_all(datetime(2026, 1, 1)) == _build_all(fixed)
    # an aware datetime in another zone is the same instant
    from datetime import timedelta, timezone

    assert _build_all(
        datetime(2026, 1, 1, 5, 30, tzinfo=timezone(timedelta(hours=5, minutes=30)))
    ) == _build_all(fixed)


@pytest.mark.parametrize("kind", ["docx", "xlsx", "pptx", "pdf"])
def test_created_changes_the_bytes_for_formats_that_embed_it(kind: str) -> None:
    assert _build_all(CREATED)[kind] != _build_all(None)[kind]


@pytest.mark.parametrize("kind", ["docx", "xlsx", "pptx"])
def test_zip_entries_are_deflated_with_pinned_timestamps(kind: str) -> None:
    for created, expected in ((None, (2026, 1, 1, 0, 0, 0)), (CREATED, (2025, 3, 14, 9, 26, 52))):
        with zipfile.ZipFile(io.BytesIO(_build_all(created)[kind])) as zf:
            assert zf.testzip() is None
            infos = zf.infolist()
            assert infos
            assert all(i.compress_type == zipfile.ZIP_DEFLATED for i in infos)
            assert all(i.date_time == expected for i in infos)


def test_zip_timestamps_clamp_to_the_zip_range() -> None:
    old = D.make_docx("T", ["p"], created=datetime(1970, 1, 1, tzinfo=UTC))
    with zipfile.ZipFile(io.BytesIO(old)) as zf:
        assert {i.date_time for i in zf.infolist()} == {(1980, 1, 1, 0, 0, 0)}
    assert _open_docx(old).paragraphs[1].text == "p"


# --------------------------------------------------------------------------------------------------
# render()
# --------------------------------------------------------------------------------------------------


@pytest.mark.parametrize("kind", ALL_KINDS)
def test_render_returns_content_and_matching_mime(kind: str) -> None:
    content, mime = D.render(kind, "Render Title", PARAGRAPHS, author="Ada", created=CREATED)
    assert isinstance(content, bytes) and content
    assert mime == D.EXTENSION_TO_MIME[kind]
    assert D.render(kind, "Render Title", PARAGRAPHS, author="Ada", created=CREATED) == (content, mime)


def test_render_content_is_readable_by_the_matching_parser() -> None:
    title = "Render Title"
    doc = _open_docx(D.render("docx", title, PARAGRAPHS, author="Ada", created=CREATED)[0])
    assert [p.text for p in doc.paragraphs] == [title, *PARAGRAPHS]
    assert doc.core_properties.author == "Ada"

    pdf = pypdf.PdfReader(io.BytesIO(D.render("pdf", title, PARAGRAPHS, author="Ada", created=CREATED)[0]))
    assert all(_squash(p) in _pdf_text(pdf) for p in PARAGRAPHS)
    assert pdf.metadata is not None and pdf.metadata.author == "Ada"

    md, _ = D.render("md", title, PARAGRAPHS)
    assert md.decode("utf-8").startswith(f"# {title}\n\n")
    txt, _ = D.render("txt", title, PARAGRAPHS)
    assert txt.decode("utf-8").startswith(f"{title}\n\n")


def test_render_xlsx_and_csv_use_table_when_given_else_single_column() -> None:
    table: list[list[object]] = [["k", "v"], ["a", 1], ["b", 2.5]]
    ws = openpyxl.load_workbook(io.BytesIO(D.render("xlsx", "T", PARAGRAPHS, table=table)[0])).active
    assert [[c.value for c in row] for row in ws.iter_rows()] == table
    assert list(csv.reader(io.StringIO(D.render("csv", "T", PARAGRAPHS, table=table)[0].decode()))) == [
        ["k", "v"],
        ["a", "1"],
        ["b", "2.5"],
    ]

    ws = openpyxl.load_workbook(io.BytesIO(D.render("xlsx", "T", PARAGRAPHS)[0])).active
    assert [row[0].value for row in ws.iter_rows()] == PARAGRAPHS
    assert max(len(list(row)) for row in ws.iter_rows()) == 1
    rows = list(csv.reader(io.StringIO(D.render("csv", "T", PARAGRAPHS)[0].decode())))
    assert rows == [[p] for p in PARAGRAPHS]


def test_render_pptx_chunks_paragraphs_into_slides_of_four_bullets() -> None:
    paragraphs = [f"bullet {i}" for i in range(9)]
    prs = pptx.Presentation(
        io.BytesIO(D.render("pptx", "Deck", paragraphs, author="Ada", created=CREATED)[0])
    )
    assert len(prs.slides) == 1 + 3
    assert prs.slides[0].shapes.title.text == "Deck"
    titles = [s.shapes.title.text for s in list(prs.slides)[1:]]
    assert titles == ["Overview", "Details", "Details 2"]
    bullets = [[p.text for p in s.placeholders[1].text_frame.paragraphs] for s in list(prs.slides)[1:]]
    assert [len(b) for b in bullets] == [4, 4, 1]
    assert [t for b in bullets for t in b] == paragraphs


def test_render_pptx_without_paragraphs_is_just_the_title_slide() -> None:
    prs = pptx.Presentation(io.BytesIO(D.render("pptx", "Deck", [])[0]))
    assert len(prs.slides) == 1


@pytest.mark.parametrize("kind", [".DOCX", " Pdf ", "markdown", "Text"])
def test_render_normalises_kind(kind: str) -> None:
    content, mime = D.render(kind, "T", ["p"])
    assert content
    assert mime in {D.MIME_DOCX, D.MIME_PDF, D.MIME_MARKDOWN, D.MIME_TEXT}


def test_render_png_is_deterministic_per_title() -> None:
    a, mime = D.render("png", "Logo A", [])
    assert mime == D.MIME_PNG
    assert a == D.render("png", "Logo A", [])[0]
    assert a != D.render("png", "Logo B", [])[0]
    assert a[:8] == b"\x89PNG\r\n\x1a\n"


def test_render_rejects_unknown_kind() -> None:
    with pytest.raises(ValueError, match="unsupported document kind"):
        D.render("exe", "T", [])
